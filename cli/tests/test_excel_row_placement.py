"""Verify actual worksheet positions instead of inferred DataFrame row counts."""

from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from openbb_cli.controllers.utils import save_to_excel


def test_new_sheet_honors_start_row(tmp_path):
    target = tmp_path / "rows.xlsx"
    save_to_excel(
        pd.DataFrame({"value": [42]}), target, "Data", start_row=3, index=False
    )
    workbook = load_workbook(target)
    try:
        assert workbook["Data"]["A4"].value == "value"
        assert workbook["Data"]["A5"].value == 42
        assert workbook["Data"]["A1"].value is None
    finally:
        workbook.close()


def test_append_uses_last_worksheet_row(tmp_path):
    target = tmp_path / "append.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "existing"
    sheet["A4"] = "footer"
    sheet["A7"].font = Font(bold=True)
    workbook.save(target)
    workbook.close()
    with patch("builtins.input", return_value="a"):
        save_to_excel(
            pd.DataFrame({"value": [42]}), target, "Data", index=False, header=False
        )
    workbook = load_workbook(target)
    try:
        assert workbook["Data"]["A4"].value == "footer"
        assert workbook["Data"]["A8"].value == 42
        assert workbook["Data"]["A7"].font.bold
    finally:
        workbook.close()
