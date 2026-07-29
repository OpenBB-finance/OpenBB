"""Overnight index fixing histories from their administrators."""

from datetime import date as dateType

NYFED_API_URL = "https://markets.newyorkfed.org/api/rates"
BOE_IADB_URL = (
    "https://www.bankofengland.co.uk/boeapps/database/_iadb-FromShowColumns.asp"
)
BOC_VALET_URL = "https://www.bankofcanada.ca/valet/observations/group/CORRA/json"
ECB_ESTR_URL = "https://data-api.ecb.europa.eu/service/data/EST/B.EU000A2X2A25.WT"
BOJ_CGI_URL = "https://www.stat-search.boj.or.jp/ssi/cgi-bin/famecgi2"
BOJ_TONA_CODE = "FM01'STRDCLUCON"
BANXICO_SERIES_URL = (
    "https://www.banxico.org.mx/SieInternet/consultarDirectorioInternetAction.do"
    "?accion=consultarSeries"
)
BANXICO_TIIE_SERIES = "SF331451"
BCB_SGS_URL = "https://api.bcb.gov.br/dados/serie/bcdata.sgs"
BCB_CDI_SERIES = "12"
SNB_DATA_PORTAL_URL = "https://data.snb.ch/api/cube/zirepo/data/json/en"
SNB_SARON_KEY = "{H0}"
MAS_DIR_URL = "https://eservices.mas.gov.sg/statistics/dir/domesticinterestrates.aspx"
MAS_SORA_COLUMN = "13"
FBIL_URL = "https://www.fbil.org.in/wasdm/ovnmibor/fetchgraphdata"
FBIL_MAX_WINDOW = "12M"
BANREP_URL = (
    "https://suameca.banrep.gov.co/estadisticas-economicas-back"
    "/rest/estadisticaEconomicaRestService/consultaInformacionSerie"
)
BANREP_IBR_SERIES = "241"
SARB_RATEREFORM_URL = "https://www.resbank.co.za/bin/sarb/ratereform"
SARB_TIMESERIES_URL = (
    "https://custom.resbank.co.za/SarbWebApi/WebIndicators/Shared"
    "/GetTimeseriesObservations"
)
SARB_JIBAR_SERIES = "MMRD403A"
BCCH_TIB_URL = "https://si3.bcentral.cl/estadisticas/Inmediato1/SerieTIB.xls"
BCCH_TIB_SHEET = "SERIE TIB"
THAIBMA_THOR_URL = "https://www.thaibma.or.th/api/thor/download-all"
BOI_SDMX_URL = (
    "https://edge.boi.gov.il/FusionEdgeServer/sdmx/v2/data/dataflow"
    "/BOI.STATISTICS/BR/1.0/"
)
BOI_SHIR_SERIES = "MNT_SHIR_D"
RBNZ_B2_URL = (
    "https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics"
    "/series/b/b2/hb2-daily-close.xlsx"
)
RBNZ_NZIONA_SERIES = "INM.DN.NZK"
RBNZ_BKBM_SERIES = "INM.DB03.NZZV"
BBK_DOWNLOAD_URL = "https://api.statistiken.bundesbank.de/rest/download"
BBK_EURIBOR_SERIES = "BBIG1/D.D0.EUR.MMKT.EURIBOR.M03.BID._Z"
CNB_PRIBOR_URL = (
    "https://www.cnb.cz/en/financial-markets/money-market/pribor"
    "/fixing-of-interest-rates-on-interbank-deposits-pribor/year.txt"
)
CNB_PRIBOR_COLUMN = "12"
CNB_MAX_YEARS = 10
BNM_MARKETS_URL = "https://financialmarkets.bnm.gov.my"
BNM_MYOR_TABLE = "data-download-myor?date_range=all_trade|2"
BNM_KLIBOR_TABLE = "data-download-klibor?date_range=all_date&date_select=&date_end=|3"
RBA_F01D_URL = "https://www.rba.gov.au/statistics/tables/xls/f01d.xlsx"
RBA_AONIA_SERIES = "FIRMMCRID"
RBA_BBSW_SERIES = "FIRMMBAB90D"
HKMA_LIQUIDITY_URL = (
    "https://api.hkma.gov.hk/public/market-data-and-statistics"
    "/daily-monetary-statistics/daily-figures-interbank-liquidity"
)
HKMA_HIBOR_FIELD = "hibor_fixing_1m"
HKMA_PAGE_SIZE = 100
BOK_ECOS_URL = "https://ecos.bok.or.kr/serviceEndpoint/httpService/request.json"
BOK_ECOS_TRX = "OSUUA02R01"
BOK_ECOS_ITEM_GROUP = "G11950"
BOK_ECOS_PAGE = 100000
GPW_DELAYED_URL = "https://gpwbenchmark.pl/delayed_data"
GPW_POLSTR_COLUMN = "1"
GPW_MAX_ATTEMPTS = 10
GPW_RETRY_SECONDS = 3.0
FINWIRE_SERIES_URL = "https://public-api.finwire.pl/v1/series/interest-rate"
FINWIRE_WIBOR_INDICATOR = "wibor-3m"
FINWIRE_MAX_DAYS = 36500
CFETS_FRR_URL = (
    "https://www.chinamoney.com.cn/r/cms/www/chinamoney/data/currency/frr-chrt.csv"
)
CFETS_FR007_COLUMN = "7"
CFETS_SHIBOR_URL = "https://www.chinamoney.com.cn/ags/ms/cm-u-bk-shibor/ShiborChrt"
CFETS_SHIBOR_TENOR = "O/N"
TAIBOR_BASE_URL = "https://www.ba.org.tw/taiborDaily"
TAIBOR_TENOR_COLUMN = "5"
TAIBOR_MAX_ARCHIVES = 40
MNB_BUBOR_URL = "https://www.mnb.hu/letoltes/bubor2.xls"
MNB_BUBOR_COLUMN = "6"
MNB_DATA_ROW = 2
BOK_CD91_SERIES = "817Y002/010502000"
BOK_KOFR_SERIES = "817Y002/010901000"

FIXING_SOURCES: dict[str, tuple[str, str]] = {
    "SOFR": ("nyfed", "secured/sofr"),
    "EFFR": ("nyfed", "unsecured/effr"),
    "CORRA": ("boc", ""),
    "ESTR": ("ecb", ""),
    "SONIA": ("boe", "IUDSOIA"),
    "TONA": ("boj", BOJ_TONA_CODE),
    "TIIE": ("banxico", BANXICO_TIIE_SERIES),
    "CDI": ("bcb", BCB_CDI_SERIES),
    "SARON": ("snb", SNB_SARON_KEY),
    "SORA": ("mas", MAS_SORA_COLUMN),
    "MIBOR": ("fbil", FBIL_MAX_WINDOW),
    "IBR": ("banrep", BANREP_IBR_SERIES),
    "ZARONIA": ("sarb", "ZARONIA"),
    "ICP": ("bcch", BCCH_TIB_SHEET),
    "THOR": ("thaibma", "O/N"),
    "SHIR": ("boi", BOI_SHIR_SERIES),
    "NZIONA": ("rbnz", RBNZ_NZIONA_SERIES),
    "BKBM": ("rbnz", RBNZ_BKBM_SERIES),
    "AONIA": ("rba", RBA_AONIA_SERIES),
    "EURIBOR": ("bbk", BBK_EURIBOR_SERIES),
    "PRIBOR": ("cnb", CNB_PRIBOR_COLUMN),
    "MYOR": ("bnm", BNM_MYOR_TABLE),
    "KLIBOR": ("bnm", BNM_KLIBOR_TABLE),
    "BBSW": ("rba", RBA_BBSW_SERIES),
    "HIBOR": ("hkma", HKMA_HIBOR_FIELD),
    "JIBAR": ("sarb_series", SARB_JIBAR_SERIES),
    "KRWCD": ("bok", BOK_CD91_SERIES),
    "KOFR": ("bok", BOK_KOFR_SERIES),
    "BUBOR": ("mnb", MNB_BUBOR_COLUMN),
    "POLSTR": ("gpw", GPW_POLSTR_COLUMN),
    "WIBOR": ("finwire", FINWIRE_WIBOR_INDICATOR),
    "FR007": ("cfets", CFETS_FR007_COLUMN),
    "SHIBOR": ("shibor", CFETS_SHIBOR_TENOR),
    "SHIBOR3M": ("shibor", "3M"),
    "TAIBOR": ("taibor", TAIBOR_TENOR_COLUMN),
}

FIXING_BASES: dict[str, float] = {
    "SOFR": 360.0,
    "EFFR": 360.0,
    "CORRA": 365.0,
    "ESTR": 360.0,
    "SONIA": 365.0,
    "TONA": 365.0,
    "TIIE": 360.0,
    "CDI": 252.0,
    "SARON": 360.0,
    "SORA": 365.0,
    "MIBOR": 365.0,
    "IBR": 360.0,
    "ZARONIA": 365.0,
    "ICP": 360.0,
    "THOR": 365.0,
    "SHIR": 365.0,
    "NZIONA": 365.0,
    "BKBM": 365.0,
    "AONIA": 365.0,
    "EURIBOR": 360.0,
    "PRIBOR": 360.0,
    "MYOR": 365.0,
    "KLIBOR": 365.0,
    "BBSW": 365.0,
    "HIBOR": 365.0,
    "JIBAR": 365.0,
    "KRWCD": 365.0,
    "KOFR": 365.0,
    "BUBOR": 360.0,
    "POLSTR": 365.0,
    "WIBOR": 365.0,
    "FR007": 365.0,
    "SHIBOR": 360.0,
    "SHIBOR3M": 360.0,
    "TAIBOR": 365.0,
}


FIXING_PROFILES: dict[str, tuple[str, str]] = {
    "SOFR": ("Secured Overnight Financing Rate", "USA"),
    "EFFR": ("Effective Federal Funds Rate", "USA"),
    "CORRA": ("Canadian Overnight Repo Rate Average", "CAN"),
    "ESTR": ("Euro Short-Term Rate", "EMU"),
    "EURIBOR": ("Euro Interbank Offered Rate (3M)", "EMU"),
    "SONIA": ("Sterling Overnight Index Average", "GBR"),
    "TONA": ("Tokyo Overnight Average Rate", "JPN"),
    "SARON": ("Swiss Average Rate Overnight", "CHE"),
    "PRIBOR": ("Prague Interbank Offered Rate (3M)", "CZE"),
    "BUBOR": ("Budapest Interbank Offered Rate (3M)", "HUN"),
    "WIBOR": ("Warsaw Interbank Offered Rate (3M)", "POL"),
    "POLSTR": ("Polish Short-Term Rate", "POL"),
    "TIIE": ("Interbank Equilibrium Interest Rate", "MEX"),
    "CDI": ("Interbank Deposit Certificate Rate", "BRA"),
    "ICP": ("Indice Camara Promedio", "CHL"),
    "IBR": ("Indicador Bancario de Referencia", "COL"),
    "AONIA": ("AUD Overnight Index Average", "AUS"),
    "BBSW": ("Bank Bill Swap Rate (3M)", "AUS"),
    "NZIONA": ("New Zealand Overnight Interbank Average", "NZL"),
    "BKBM": ("Bank Bill Benchmark Rate (3M)", "NZL"),
    "SORA": ("Singapore Overnight Rate Average", "SGP"),
    "MYOR": ("Malaysia Overnight Rate", "MYS"),
    "KLIBOR": ("Kuala Lumpur Interbank Offered Rate (3M)", "MYS"),
    "THOR": ("Thai Overnight Repurchase Rate", "THA"),
    "HIBOR": ("Hong Kong Interbank Offered Rate (1M)", "HKG"),
    "FR007": ("7-Day Fixing Repo Rate", "CHN"),
    "SHIBOR": ("Shanghai Interbank Offered Rate (O/N)", "CHN"),
    "SHIBOR3M": ("Shanghai Interbank Offered Rate (3M)", "CHN"),
    "TAIBOR": ("Taipei Interbank Offered Rate (3M)", "TWN"),
    "KOFR": ("Korea Overnight Financing Repo Rate", "KOR"),
    "KRWCD": ("Certificate of Deposit Rate (91-day)", "KOR"),
    "MIBOR": ("Mumbai Interbank Outright Rate", "IND"),
    "SHIR": ("Shekel Interbank Rate", "ISR"),
    "ZARONIA": ("South African Rand Overnight Index Average", "ZAF"),
    "JIBAR": ("Johannesburg Interbank Average Rate (3M)", "ZAF"),
}


def fixing_profile(index: str) -> tuple[str, str]:
    """Readable name and country ISO3 an index is published for."""
    return FIXING_PROFILES.get(index, (index, ""))


OIS_CURRENCY_INDICES: dict[str, str] = {
    "USD": "SOFR",
    "GBP": "SONIA",
    "JPY": "TONA",
    "EUR": "ESTR",
    "CHF": "SARON",
    "CAD": "CORRA",
    "SGD": "SORA",
    "ZAR": "ZARONIA",
    "ILS": "SHIR",
    "AUD": "AONIA",
    "NZD": "NZIONA",
    "BRL": "CDI",
    "INR": "MIBOR",
    "COP": "IBR",
    "CLP": "ICP",
    "THB": "THOR",
    "MXN": "TIIE",
    "MYR": "MYOR",
    "KRW": "KOFR",
    "PLN": "POLSTR",
    "CNY": "FR007",
}

FIXED_FLOAT_CURRENCY_INDICES: dict[str, str] = {
    "EUR": "EURIBOR",
    "CZK": "PRIBOR",
    "USD": "SOFR",
    "GBP": "SONIA",
    "JPY": "TONA",
    "CHF": "SARON",
    "CAD": "CORRA",
    "NZD": "BKBM",
    "SGD": "SORA",
    "BRL": "CDI",
    "COP": "IBR",
    "CLP": "ICP",
    "THB": "THOR",
    "INR": "MIBOR",
    "MXN": "TIIE",
    "MYR": "KLIBOR",
    "AUD": "BBSW",
    "HKD": "HIBOR",
    "ZAR": "JIBAR",
    "KRW": "KRWCD",
    "HUF": "BUBOR",
    "PLN": "WIBOR",
    "CNY": "FR007",
    "TWD": "TAIBOR",
}


def _currency_scoped_index(name: str) -> str | None:
    """Index a currency-scoped FISN underlier resolves to, by swap family."""
    for prefix, table in (
        ("NA/SWAP OIS ", OIS_CURRENCY_INDICES),
        ("NA/SWAP FXD FLT ", FIXED_FLOAT_CURRENCY_INDICES),
    ):
        if name.startswith(prefix):
            tokens = name[len(prefix) :].split()

            if len(tokens) == 1:
                return table.get(tokens[0])

    return None


def index_for_underlier(underlier: str | None) -> str | None:
    """Return the fixing index an underlier name compounds, or None when unsupported."""
    name = (underlier or "").upper()

    if "SOFR" in name:
        return "SOFR"

    if "FEDERAL FUNDS" in name:
        return "EFFR"

    if "CORRA" in name:
        return "CORRA"

    if "EUROSTR" in name:
        return "ESTR"

    if "EURIBOR" in name and "SWAP RATE" not in name:
        return "EURIBOR"

    if "PRIBOR" in name:
        return "PRIBOR"

    if "MYOR" in name:
        return "MYOR"

    if "KLIBOR" in name:
        return "KLIBOR"

    if "SONIA" in name:
        return "SONIA"

    if "TONA" in name:
        return "TONA"

    if "TIIE" in name:
        return "TIIE"

    if "CDI" in name:
        return "CDI"

    if "SARON" in name:
        return "SARON"

    if "SORA" in name:
        return "SORA"

    if "MIBOR" in name:
        return "MIBOR"

    if "IBR" in name:
        return "IBR"

    if "ZARONIA" in name:
        return "ZARONIA"

    if "ICP" in name or "CLP-TNA" in name:
        return "ICP"

    if "THOR" in name:
        return "THOR"

    if "SHIR" in name:
        return "SHIR"

    if "NZIONA" in name:
        return "NZIONA"

    if "BKBM" in name or "BBR" in name:
        return "BKBM"

    if "AONIA" in name:
        return "AONIA"

    if "BBSW" in name or "BBSY" in name:
        return "BBSW"

    if "HIBOR" in name:
        return "HIBOR"

    if "JIBAR" in name:
        return "JIBAR"

    if "KOFR" in name:
        return "KOFR"

    if "BUBOR" in name:
        return "BUBOR"

    if "POLSTR" in name or "WIRON" in name:
        return "POLSTR"

    if "WIBOR" in name:
        return "WIBOR"

    if "FIXING REPO" in name or "CNREPOFIX" in name or "FR007" in name:
        return "FR007"

    if "TAIBOR" in name:
        return "TAIBOR"

    if "KRW-CD" in name or "KRW CD" in name:
        return "KRWCD"

    return _currency_scoped_index(name)


def fixing_basis(index: str | None) -> float:
    """Accrual basis an index's compounding is defined on."""
    return FIXING_BASES.get(index or "", 360.0)


async def _fetch_nyfed(path: str, start: str, end: str) -> dict[str, float]:
    """Fetch a New York Fed reference rate series over a window."""
    from openbb_core.app.model.abstract.error import OpenBBError
    from openbb_core.provider.utils.helpers import amake_request

    url = (
        f"{NYFED_API_URL}/{path}/search.json?type=rate&startDate={start}&endDate={end}"
    )
    response = await amake_request(url)
    rows = (response or {}).get("refRates") if isinstance(response, dict) else None

    if rows is None:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    return {
        (row.get("effectiveDate") or "").strip(): float(row["percentRate"]) / 100.0
        for row in rows
        if row.get("percentRate") is not None
    }


async def _fetch_corra(start: str, end: str) -> dict[str, float]:
    """Fetch CORRA from the Bank of Canada Valet API over a window."""
    from openbb_core.app.model.abstract.error import OpenBBError
    from openbb_core.provider.utils.helpers import amake_request

    url = f"{BOC_VALET_URL}?start_date={start}&end_date={end}"
    response = await amake_request(url)
    rows = (response or {}).get("observations") if isinstance(response, dict) else None

    if rows is None:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    fixings: dict[str, float] = {}

    for row in rows:
        day = (row.get("d") or "").strip()
        value = (row.get("AVG.INTWO") or {}).get("v")

        if day and value is not None:
            fixings[day] = float(value) / 100.0

    return fixings


async def _fetch_estr(start: str, end: str) -> dict[str, float]:
    """Fetch the euro short-term rate from the ECB data API, served as CSV."""
    import csv
    import io

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    url = f"{ECB_ESTR_URL}?format=csvdata&startPeriod={start}&endPeriod={end}"

    async with (
        aiohttp.ClientSession() as session,
        session.get(url) as response,
    ):
        response.raise_for_status()
        payload = await response.text()

    if not isinstance(payload, str) or "TIME_PERIOD" not in payload:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    fixings: dict[str, float] = {}

    for row in csv.DictReader(io.StringIO(payload)):
        day = (row.get("TIME_PERIOD") or "").strip()
        value = (row.get("OBS_VALUE") or "").strip()

        if day and value:
            fixings[day] = float(value) / 100.0

    return fixings


async def _fetch_boe(series: str, start: str, end: str) -> dict[str, float]:
    """Fetch a Bank of England IADB series over a window, served as CSV."""
    from datetime import datetime

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    from openbb_cftc.utils.search import SEARCH_USER_AGENT

    def _bound(day: str) -> str:
        return datetime.strptime(day, "%Y-%m-%d").strftime("%d/%b/%Y")

    url = (
        f"{BOE_IADB_URL}?csv.x=yes&Datefrom={_bound(start)}&Dateto={_bound(end)}"
        f"&SeriesCodes={series}&CSVF=TN&UsingCodes=Y&VPD=Y&VFD=N"
    )

    async with (
        aiohttp.ClientSession(headers={"User-Agent": SEARCH_USER_AGENT}) as session,
        session.get(url) as response,
    ):
        response.raise_for_status()
        payload = await response.text()

    if "DATE," not in payload:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    fixings: dict[str, float] = {}

    for line in payload.splitlines()[1:]:
        parts = line.split(",")

        if len(parts) == 2 and parts[1].strip():
            day = datetime.strptime(parts[0].strip(), "%d %b %Y").date().isoformat()
            fixings[day] = float(parts[1]) / 100.0

    return fixings


async def _fetch_boj(series: str, start: str, end: str) -> dict[str, float]:
    """Fetch a Bank of Japan series via the stat-search CGI's two-step POST/GET CSV flow."""
    import re
    from datetime import datetime

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    from openbb_cftc.utils.search import SEARCH_USER_AGENT

    year_from = datetime.strptime(start, "%Y-%m-%d").year
    year_to = datetime.strptime(end, "%Y-%m-%d").year
    headers = {"User-Agent": SEARCH_USER_AGENT}

    async with aiohttp.ClientSession(headers=headers) as session:
        async with session.post(
            BOJ_CGI_URL,
            data={
                "cgi": "$nme_r000_en",
                "obj_name": series.split("'", maxsplit=1)[0],
                "txtYyyyFrom": str(year_from),
                "txtYyyyTo": str(year_to),
                "cmbFREQ": "NONE",
                "cmbFREQ_OPTION": "NONE",
                "lstCode": series,
            },
        ) as response:
            response.raise_for_status()
            result_page = await response.text()

        if "DLform_DD" not in result_page:
            raise OpenBBError(f"Unexpected fixings response from {BOJ_CGI_URL}.")

        async with session.post(
            BOJ_CGI_URL,
            data={
                "cgi": "$nme_r030_en",
                "chkfrq": "DD",
                "rdoheader": "SIMPLE",
                "rdodelimitar": "COMMA",
                "hdnYyyyFrom": str(year_from),
                "hdnYyyyTo": str(year_to),
                "hdncode": series,
                "sw_freq": "NONE",
                "sw_yearend": "NONE",
                "sw_observed": "NONE",
            },
        ) as response:
            response.raise_for_status()
            download_page = await response.text()

        match = re.search(r'href="(/ssi/html/[^"]+\.csv)"', download_page)

        if match is None:
            raise OpenBBError(f"Unexpected fixings response from {BOJ_CGI_URL}.")

        csv_url = "https://www.stat-search.boj.or.jp" + match.group(1)

        async with session.get(csv_url) as response:
            response.raise_for_status()
            payload = await response.text()

    fixings: dict[str, float] = {}

    for line in payload.splitlines():
        parts = line.split(",", 1)

        if len(parts) != 2 or not re.match(r"^\d{4}/\d{2}/\d{2}$", parts[0]):
            continue

        value = parts[1].strip()

        if value and value != "NA":
            day = datetime.strptime(parts[0], "%Y/%m/%d").date().isoformat()
            fixings[day] = float(value) / 100.0

    return fixings


async def _fetch_banxico(series: str, start: str, end: str) -> dict[str, float]:
    """Fetch a Banco de Mexico SIE series export over a window, served as CSV."""
    import re
    from datetime import datetime

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    year_from = datetime.strptime(start, "%Y-%m-%d").year
    year_to = datetime.strptime(end, "%Y-%m-%d").year

    async with (
        aiohttp.ClientSession() as session,
        session.post(
            BANXICO_SERIES_URL,
            data={
                "locale": "en",
                "idCuadro": "CA51",
                "sector": "18",
                "version": "3",
                "series": series,
                "anoInicial": str(year_from),
                "anoFinal": str(year_to),
                "tipoInformacion": "4,1",
                "formatoHorizontal": "false",
                "metadatosWeb": "false",
                "formatoCSV.x": "1",
                "formatoCSV.y": "1",
            },
        ) as response,
    ):
        response.raise_for_status()
        payload = await response.text()

    if series not in payload:
        raise OpenBBError(f"Unexpected fixings response from {BANXICO_SERIES_URL}.")

    fixings: dict[str, float] = {}

    for line in payload.splitlines():
        parts = line.split(",")

        if len(parts) != 2 or not re.match(r"^\d{2}/\d{2}/\d{4}$", parts[0]):
            continue

        value = parts[1].strip()

        if value:
            day = datetime.strptime(parts[0], "%m/%d/%Y").date().isoformat()
            fixings[day] = float(value) / 100.0

    return fixings


async def _fetch_bcb(series: str, start: str, end: str) -> dict[str, float]:
    """Fetch a Banco Central do Brasil SGS series over a window, served as JSON."""
    from datetime import datetime

    from openbb_core.app.model.abstract.error import OpenBBError
    from openbb_core.provider.utils.helpers import amake_request

    def _bound(day: str) -> str:
        return datetime.strptime(day, "%Y-%m-%d").strftime("%d/%m/%Y")

    url = (
        f"{BCB_SGS_URL}.{series}/dados?formato=json"
        f"&dataInicial={_bound(start)}&dataFinal={_bound(end)}"
    )
    response = await amake_request(url)

    if not isinstance(response, list):
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    fixings: dict[str, float] = {}

    for row in response:
        day = (row.get("data") or "").strip()
        value = row.get("valor")

        if day and value is not None:
            fixings[datetime.strptime(day, "%d/%m/%Y").date().isoformat()] = (
                float(value) / 100.0
            )

    return fixings


async def _fetch_snb(key: str, start: str, end: str) -> dict[str, float]:
    """Fetch an SNB Data Portal repo-reference-rate series over a window, as JSON."""
    from openbb_core.app.model.abstract.error import OpenBBError
    from openbb_core.provider.utils.helpers import amake_request

    url = f"{SNB_DATA_PORTAL_URL}?fromDate={start}&toDate={end}"
    response = await amake_request(url)
    series = (response or {}).get("timeseries") if isinstance(response, dict) else None

    if series is None:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    wanted = next(
        (s for s in series if s.get("metadata", {}).get("key", "").endswith(key)), None
    )

    if wanted is None:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    return {
        row["date"]: float(row["value"]) / 100.0
        for row in wanted.get("values", [])
        if row.get("value") is not None
    }


async def _fetch_mas(column: str, start: str, end: str) -> dict[str, float]:
    """Fetch SORA from MAS's Domestic Interest Rates page via its GET/POST ASP.NET postback."""
    import re
    from datetime import datetime

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    from openbb_cftc.utils.search import SEARCH_USER_AGENT

    year_from = datetime.strptime(start, "%Y-%m-%d").year
    year_to = datetime.strptime(end, "%Y-%m-%d").year
    headers = {"User-Agent": SEARCH_USER_AGENT}

    async with aiohttp.ClientSession(headers=headers) as session:
        async with session.get(MAS_DIR_URL) as response:
            response.raise_for_status()
            page = await response.text()

        matches = {
            name: re.search(rf'id="{name}" value="([^"]*)"', page)
            for name in ("__VIEWSTATE", "__VIEWSTATEGENERATOR", "__EVENTVALIDATION")
        }

        if any(match is None for match in matches.values()):
            raise OpenBBError(f"Unexpected fixings response from {MAS_DIR_URL}.")

        tokens = {name: match.group(1) for name, match in matches.items() if match}

        async with session.post(
            MAS_DIR_URL,
            data={
                **tokens,
                "ctl00$ContentPlaceHolder1$StartYearDropDownList": str(year_from),
                "ctl00$ContentPlaceHolder1$EndYearDropDownList": str(year_to),
                "ctl00$ContentPlaceHolder1$StartMonthDropDownList": "1",
                "ctl00$ContentPlaceHolder1$EndMonthDropDownList": "12",
                f"ctl00$ContentPlaceHolder1$ColumnsCheckBoxList${column}": "on",
                "ctl00$ContentPlaceHolder1$Button1": "Display",
            },
        ) as response:
            response.raise_for_status()
            result_page = await response.text()

    if "SORA VALUE" not in result_page:
        raise OpenBBError(f"Unexpected fixings response from {MAS_DIR_URL}.")

    fixings: dict[str, float] = {}
    year = month = None

    def _cell(raw: str) -> str:
        stripped = raw.strip()

        return "" if stripped == "&nbsp;" else stripped

    for row in re.finditer(r"<tr>(.*?)</tr>", result_page, re.S):
        cells = re.findall(r"<td[^>]*>(.*?)</td>", row.group(1), re.S)

        if len(cells) != 5:
            continue

        year_cell, month_cell, day_cell, _, value_cell = (_cell(c) for c in cells)
        year = year_cell or year
        month = month_cell or month

        if not (year and month and day_cell and value_cell):
            continue

        day = datetime.strptime(f"{year}-{month}-{day_cell}", "%Y-%b-%d").date()
        fixings[day.isoformat()] = float(value_cell) / 100.0

    return fixings


async def _fetch_fbil(window: str, start: str, end: str) -> dict[str, float]:
    """Fetch Overnight MIBOR from FBIL's chart-data endpoint."""
    from openbb_core.app.model.abstract.error import OpenBBError
    from openbb_core.provider.utils.helpers import amake_request

    url = f"{FBIL_URL}/{window}"
    response = await amake_request(url)
    trend = (
        (response or {}).get("benchMarkTrend") if isinstance(response, dict) else None
    )

    if trend is None:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    return {
        row["publishDate"]: float(row["rate"]) / 100.0
        for row in trend
        if row.get("publishDate") and start <= row["publishDate"] <= end
    }


async def _fetch_banrep(series: str, start: str, end: str) -> dict[str, float]:
    """Fetch Banco de la Republica's IBR series."""
    from datetime import datetime, timezone

    from openbb_core.app.model.abstract.error import OpenBBError
    from openbb_core.provider.utils.helpers import amake_request

    url = f"{BANREP_URL}?idSerie={series}"
    response = await amake_request(url)
    series_list = response if isinstance(response, list) else None

    if not series_list or not isinstance(series_list[0], dict):
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    points = series_list[0].get("data")

    if points is None:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    fixings: dict[str, float] = {}

    for point in points:
        millis, value = point
        day = datetime.fromtimestamp(millis / 1000, tz=timezone.utc).date().isoformat()

        if start <= day <= end:
            fixings[day] = float(value) / 100.0

    return fixings


async def _fetch_sarb(rate_type: str, start: str, end: str) -> dict[str, float]:
    """Fetch ZARONIA from SARB's rate-reform API."""
    import json

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    url = (
        f"{SARB_RATEREFORM_URL}?start_date={start}&end_date={end}&rate_type={rate_type}"
    )

    async with (
        aiohttp.ClientSession() as session,
        session.get(url) as response,
    ):
        response.raise_for_status()
        payload = await response.text()

    try:
        parsed = json.loads(payload)
    except ValueError:
        parsed = None

    rows = (parsed or {}).get("rates") if isinstance(parsed, dict) else None

    if rows is None:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    return {
        (row.get("date") or "")[:10]: float(row["rate"]) / 100.0
        for row in rows
        if row.get("date") and row.get("rate") is not None
    }


async def _fetch_bcch(sheet: str, start: str, end: str) -> dict[str, float]:
    """Fetch Banco Central de Chile's TIB series."""
    import aiohttp
    import xlrd
    from openbb_core.app.model.abstract.error import OpenBBError

    async with (
        aiohttp.ClientSession() as session,
        session.get(BCCH_TIB_URL) as response,
    ):
        response.raise_for_status()
        content = await response.read()

    try:
        book = xlrd.open_workbook(file_contents=content)
        rows = book.sheet_by_name(sheet)
    except Exception as exc:
        raise OpenBBError(f"Unexpected fixings response from {BCCH_TIB_URL}.") from exc

    fixings: dict[str, float] = {}

    for row_index in range(rows.nrows):
        if rows.cell_type(row_index, 0) != xlrd.XL_CELL_DATE:
            continue

        day = xlrd.xldate.xldate_as_datetime(
            rows.cell_value(row_index, 0), book.datemode
        ).date()
        day_iso = day.isoformat()

        if not start <= day_iso <= end:
            continue

        try:
            rate = float(rows.cell_value(row_index, 2))
        except (TypeError, ValueError):
            continue

        fixings[day_iso] = rate / 100.0

    return fixings


async def _fetch_thaibma(tenor: str, start: str, end: str) -> dict[str, float]:
    """Fetch THOR from ThaiBMA's full-history CSV export."""
    import csv
    import io

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    async with (
        aiohttp.ClientSession() as session,
        session.get(THAIBMA_THOR_URL) as response,
    ):
        response.raise_for_status()
        payload = await response.text()

    if "As of" not in payload:
        raise OpenBBError(f"Unexpected fixings response from {THAIBMA_THOR_URL}.")

    fixings: dict[str, float] = {}

    for row in csv.DictReader(io.StringIO(payload)):
        day = (row.get("As of") or "").strip()
        code = (row.get("Code") or "").strip()
        row_tenor = (row.get("Tenor") or "").strip()
        value = (row.get("Rate") or "").strip()

        if (
            code == "THOR"
            and row_tenor == tenor
            and day
            and value
            and start <= day <= end
        ):
            fixings[day] = float(value) / 100.0

    return fixings


async def _fetch_boi(series: str, start: str, end: str) -> dict[str, float]:
    """Fetch SHIR from the Bank of Israel's SDMX v2 API."""
    from openbb_core.app.model.abstract.error import OpenBBError
    from openbb_core.provider.utils.helpers import amake_request

    url = (
        f"{BOI_SDMX_URL}?c[SERIES_CODE]={series}&startPeriod={start}&endPeriod={end}"
        "&format=sdmx-json&dimensionAtObservation=AllDimensions&locale=en"
    )
    response = await amake_request(url)

    if not isinstance(response, dict):
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    try:
        observations = response["data"]["dataSets"][0]["observations"]
        dimensions = response["data"]["structure"]["dimensions"]["observation"]
        periods = next(d for d in dimensions if d["id"] == "TIME_PERIOD")["values"]
    except (KeyError, IndexError, StopIteration, TypeError) as exc:
        raise OpenBBError(f"Unexpected fixings response from {url}.") from exc

    fixings: dict[str, float] = {}

    for key, value in observations.items():
        index = int(key.rsplit(":", maxsplit=1)[-1])
        rate = value[0] if value else None

        if rate is not None:
            fixings[periods[index]["id"]] = float(rate) / 100.0

    return fixings


def _taibor_rows(text: str, index: int, start: str, end: str) -> dict[str, float]:
    """Parse dated TAIBOR rows out of one archive sheet."""
    import re
    from datetime import datetime

    fixings: dict[str, float] = {}

    for line in text.splitlines():
        cells = line.split(",")

        if len(cells) <= index or not re.fullmatch(
            r"\d{2}[A-Z]{3}\d{2}", cells[0].strip()
        ):
            continue

        try:
            day = datetime.strptime(cells[0].strip(), "%d%b%y").date().isoformat()
            rate = float(cells[index].strip())
        except ValueError:
            continue

        if start <= day <= end:
            fixings[day] = rate / 100.0

    return fixings


async def _fetch_taibor(column: str, start: str, end: str) -> dict[str, float]:
    """Fetch TAIBOR from the Bankers Association's monthly archive downloads."""
    import io
    import re
    import ssl
    import zipfile

    import aiohttp
    import certifi
    from openbb_core.app.model.abstract.error import OpenBBError

    index = int(column)
    context = ssl.create_default_context(cafile=certifi.where())
    context.verify_flags &= ~ssl.VERIFY_X509_STRICT
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
            " (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
        )
    }
    fixings: dict[str, float] = {}

    async with aiohttp.ClientSession(
        headers=headers, connector=aiohttp.TCPConnector(ssl=context)
    ) as session:
        async with session.get(f"{TAIBOR_BASE_URL}/index") as response:
            response.raise_for_status()
            listing = await response.text()

        wanted: list[str] = []

        for row in re.findall(r"<tr[^>]*>(.*?)</tr>", listing, re.S | re.I):
            stamp = re.search(r"(\d{4})/(\d{2})/(\d{2})", row)
            link = re.search(r"DownloadAll\?id=(\d+)", row)

            if not stamp or not link:
                continue

            posted = f"{stamp.group(1)}-{stamp.group(2)}-{stamp.group(3)}"

            if posted >= start and len(wanted) < TAIBOR_MAX_ARCHIVES:
                wanted.append(link.group(1))

        for archive in wanted:
            async with session.get(
                f"{TAIBOR_BASE_URL}/DownloadAll?id={archive}"
            ) as response:
                if response.status != 200:
                    continue

                blob = await response.read()

            try:
                bundle = zipfile.ZipFile(io.BytesIO(blob))
            except zipfile.BadZipFile:
                continue

            for name in bundle.namelist():
                if not name.lower().endswith(".csv"):
                    continue

                raw = bundle.read(name)

                for encoding in ("big5", "cp950", "utf-8-sig"):
                    try:
                        fixings.update(
                            _taibor_rows(raw.decode(encoding), index, start, end)
                        )
                        break
                    except UnicodeDecodeError:
                        continue

    if not fixings:
        raise OpenBBError(f"Unexpected fixings response from {TAIBOR_BASE_URL}.")

    return fixings


async def _fetch_cfets(column: str, start: str, end: str) -> dict[str, float]:
    """Fetch a fixing repo rate column from the CFETS chart series."""
    import re
    import ssl

    import aiohttp
    import certifi
    from openbb_core.app.model.abstract.error import OpenBBError

    index = int(column)
    context = ssl.create_default_context(cafile=certifi.where())
    context.verify_flags &= ~ssl.VERIFY_X509_STRICT
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
            " (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
        ),
        "Referer": "https://www.chinamoney.com.cn/english/bmkfrr/",
    }

    async with (
        aiohttp.ClientSession(
            headers=headers, connector=aiohttp.TCPConnector(ssl=context)
        ) as session,
        session.get(CFETS_FRR_URL) as response,
    ):
        response.raise_for_status()
        payload = await response.text()

    fixings: dict[str, float] = {}

    for line in payload.splitlines():
        cells = line.split(",")

        if len(cells) <= index or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", cells[0]):
            continue

        try:
            rate = float(cells[index])
        except ValueError:
            continue

        if start <= cells[0] <= end:
            fixings[cells[0]] = rate / 100.0

    if not fixings:
        raise OpenBBError(f"Unexpected fixings response from {CFETS_FRR_URL}.")

    return fixings


async def _fetch_shibor(tenor: str, start: str, end: str) -> dict[str, float]:
    """Fetch a SHIBOR tenor from the CFETS benchmark chart service."""
    import json
    import re
    import ssl

    import aiohttp
    import certifi
    from openbb_core.app.model.abstract.error import OpenBBError

    context = ssl.create_default_context(cafile=certifi.where())
    context.verify_flags &= ~ssl.VERIFY_X509_STRICT
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
            " (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
        ),
        "Referer": "https://www.chinamoney.com.cn/english/bmkshb/",
    }

    async with (
        aiohttp.ClientSession(
            headers=headers, connector=aiohttp.TCPConnector(ssl=context)
        ) as session,
        session.post(CFETS_SHIBOR_URL) as response,
    ):
        response.raise_for_status()
        payload = json.loads(await response.text())

    data = payload.get("data") or {}
    columns = data.get("columns") or []
    fixings: dict[str, float] = {}

    if tenor in columns:
        index = columns.index(tenor)

        for line in (data.get("csv") or "").splitlines():
            cells = line.split(",")

            if len(cells) <= index or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", cells[0]):
                continue

            try:
                rate = float(cells[index])
            except ValueError:
                continue

            if start <= cells[0] <= end:
                fixings[cells[0]] = rate / 100.0

    if not fixings:
        raise OpenBBError(f"Unexpected fixings response from {CFETS_SHIBOR_URL}.")

    return fixings


async def _fetch_finwire(indicator: str, start: str, end: str) -> dict[str, float]:
    """Fetch a Polish benchmark series from the finwire public data API."""
    import json
    import ssl
    from datetime import date as _date

    import aiohttp
    import certifi
    from openbb_core.app.model.abstract.error import OpenBBError

    context = ssl.create_default_context(cafile=certifi.where())
    context.verify_flags &= ~ssl.VERIFY_X509_STRICT
    span = (_date.fromisoformat(end) - _date.fromisoformat(start)).days + 1
    days = min(max(span, 1), FINWIRE_MAX_DAYS)
    url = f"{FINWIRE_SERIES_URL}/{indicator}?days={days}"

    async with (
        aiohttp.ClientSession(connector=aiohttp.TCPConnector(ssl=context)) as session,
        session.get(url) as response,
    ):
        response.raise_for_status()
        payload = json.loads(await response.text())

    points = ((payload or {}).get("data") or {}).get("points") or []
    fixings: dict[str, float] = {}

    for point in points:
        day = (point.get("date") or "").strip()
        value = point.get("value")

        if len(day) == 10 and value is not None and start <= day <= end:
            fixings[day] = float(value) / 100.0

    if not fixings:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    return fixings


async def _fetch_gpw(column: str, start: str, end: str) -> dict[str, float]:
    """Fetch a POLSTR column from the GPW Benchmark delayed data table."""
    import re
    import ssl

    import aiohttp
    import certifi
    from openbb_core.app.model.abstract.error import OpenBBError

    index = int(column)
    context = ssl.create_default_context(cafile=certifi.where())
    context.verify_flags &= ~ssl.VERIFY_X509_STRICT
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
            " (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
        )
    }

    import asyncio

    payload = ""

    for attempt in range(GPW_MAX_ATTEMPTS):
        try:
            async with (
                aiohttp.ClientSession(
                    headers=headers, connector=aiohttp.TCPConnector(ssl=context)
                ) as session,
                session.get(GPW_DELAYED_URL) as response,
            ):
                response.raise_for_status()
                payload = await response.text()

            break
        except aiohttp.ClientError:
            if attempt == GPW_MAX_ATTEMPTS - 1:
                raise

            await asyncio.sleep(GPW_RETRY_SECONDS)

    fixings: dict[str, float] = {}

    for row in re.findall(r"<tr[^>]*>(.*?)</tr>", payload, re.S | re.I):
        cells = [
            re.sub(r"<[^>]+>", "", cell).replace("&nbsp;", " ").strip()
            for cell in re.findall(r"<td[^>]*>(.*?)</td>", row, re.S | re.I)
        ]

        if len(cells) <= index or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", cells[0]):
            continue

        try:
            rate = float(cells[index].replace(",", ""))
        except ValueError:
            continue

        if start <= cells[0] <= end:
            fixings[cells[0]] = rate / 100.0

    if not fixings:
        raise OpenBBError(f"Unexpected fixings response from {GPW_DELAYED_URL}.")

    return fixings


async def _fetch_mnb(column: str, start: str, end: str) -> dict[str, float]:
    """Fetch a BUBOR tenor from the MNB workbook's per-year fixing sheets."""
    import ssl
    from datetime import date as _date

    import aiohttp
    import certifi
    import xlrd
    from openbb_core.app.model.abstract.error import OpenBBError

    index = int(column)
    connector = aiohttp.TCPConnector(
        ssl=ssl.create_default_context(cafile=certifi.where())
    )

    async with (
        aiohttp.ClientSession(connector=connector) as session,
        session.get(MNB_BUBOR_URL) as response,
    ):
        response.raise_for_status()
        content = await response.read()

    book = xlrd.open_workbook(file_contents=content)
    fixings: dict[str, float] = {}

    for name in book.sheet_names():
        if not name.isdigit() or not start[:4] <= name <= end[:4]:
            continue

        sheet = book.sheet_by_name(name)

        for row in range(MNB_DATA_ROW, sheet.nrows):
            stamp = sheet.cell_value(row, 0)
            value = sheet.cell_value(row, index) if index < sheet.ncols else ""

            if not isinstance(stamp, float) or stamp <= 0:
                continue

            try:
                rate = float(value)
            except (TypeError, ValueError):
                continue

            parts = xlrd.xldate_as_tuple(stamp, book.datemode)
            day = _date(*parts[:3]).isoformat()

            if start <= day <= end:
                fixings[day] = rate / 100.0

    if not fixings:
        raise OpenBBError(f"Unexpected fixings response from {MNB_BUBOR_URL}.")

    return fixings


async def _fetch_bok(series: str, start: str, end: str) -> dict[str, float]:
    """Fetch a Bank of Korea ECOS daily series by statistic table and item code."""
    import json
    import re
    import ssl

    import aiohttp
    import certifi
    from openbb_core.app.model.abstract.error import OpenBBError

    table, _, item = series.partition("/")
    context = ssl.create_default_context(cafile=certifi.where())
    context.verify_flags &= ~ssl.VERIFY_X509_STRICT
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
            " (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
        ),
        "Content-Type": "application/json",
        "Referer": "https://ecos.bok.or.kr/",
    }
    request = {
        "header": {
            "guidSeq": 1,
            "trxCd": BOK_ECOS_TRX,
            "scrId": "IECOSPCM02",
            "sysCd": "03",
            "fstChnCd": "WEB",
            "langDvsnCd": "EN",
            "envDvsnCd": "D",
            "sndRspnDvsnCd": "S",
            "sndDtm": end.replace("-", ""),
            "usrId": "IECOSPC",
            "pageNum": 1,
            "pageCnt": BOK_ECOS_PAGE,
        },
        "data": {
            "statSrchDsList": [
                {
                    "dsId": table,
                    "dsItmId1": "ACC_ITEM",
                    "dsItmGrpId1": BOK_ECOS_ITEM_GROUP,
                    "dsItmVal1": item,
                }
            ],
            "statSrchFreqList": [
                {
                    "freq": "D",
                    "vlidStDtm": start.replace("-", ""),
                    "vlidEndDtm": end.replace("-", ""),
                }
            ],
            "statTyp": "E",
            "statDataCvsnCdList": ["00"],
            "viewType": "01",
            "holidayYn": "Y",
        },
    }

    async with (
        aiohttp.ClientSession(
            headers=headers, connector=aiohttp.TCPConnector(ssl=context)
        ) as session,
        session.post(BOK_ECOS_URL, json=request) as response,
    ):
        response.raise_for_status()
        payload = json.loads(await response.text())

    content = ((payload or {}).get("data") or {}).get("jsonCtnt")

    if not content:
        raise OpenBBError("Unexpected fixings response from ECOS.")

    fixings: dict[str, float] = {}

    for row in json.loads(content):
        for key, value in row.items():
            if not re.fullmatch(r"\d{8}", key) or value in (None, ""):
                continue

            try:
                rate = float(value)
            except (TypeError, ValueError):
                continue

            fixings[f"{key[:4]}-{key[4:6]}-{key[6:]}"] = rate / 100.0

    if not fixings:
        raise OpenBBError("Unexpected fixings response from ECOS.")

    return fixings


async def _fetch_sarb_series(series: str, start: str, end: str) -> dict[str, float]:
    """Fetch a SARB web-indicator timeseries by its code."""
    import json

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    url = f"{SARB_TIMESERIES_URL}/{series}"

    async with (
        aiohttp.ClientSession() as session,
        session.get(url) as response,
    ):
        response.raise_for_status()
        rows = json.loads(await response.text())

    if not isinstance(rows, list):
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    fixings: dict[str, float] = {}

    for row in rows:
        day = (row.get("Period") or "")[:10]
        value = row.get("Value")

        if day and value is not None and start <= day <= end:
            fixings[day] = float(value) / 100.0

    if not fixings:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    return fixings


async def _fetch_hkma(field: str, start: str, end: str) -> dict[str, float]:
    """Fetch a HIBOR fixing from the HKMA daily interbank liquidity series."""
    import json

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    fixings: dict[str, float] = {}
    offset = 0

    async with aiohttp.ClientSession() as session:
        while True:
            url = (
                f"{HKMA_LIQUIDITY_URL}?from={start}&to={end}"
                f"&pagesize={HKMA_PAGE_SIZE}&offset={offset}"
            )

            async with session.get(url) as response:
                response.raise_for_status()
                payload = json.loads(await response.text())

            records = ((payload or {}).get("result") or {}).get("records") or []

            for record in records:
                day = (record.get("end_of_date") or "").strip()
                value = record.get(field)

                if day and value is not None and start <= day <= end:
                    fixings[day] = float(value) / 100.0

            if len(records) < HKMA_PAGE_SIZE:
                break

            offset += HKMA_PAGE_SIZE

    if not fixings:
        raise OpenBBError(f"Unexpected fixings response from {HKMA_LIQUIDITY_URL}.")

    return fixings


async def _fetch_bnm(table: str, start: str, end: str) -> dict[str, float]:
    """Fetch a rate column from a Bank Negara financial markets table."""
    import re

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    page, _, column = table.partition("|")
    index = int(column)
    url = f"{BNM_MARKETS_URL}/{page}"

    async with (
        aiohttp.ClientSession() as session,
        session.get(url) as response,
    ):
        response.raise_for_status()
        payload = await response.text()

    fixings: dict[str, float] = {}

    for row in re.findall(r"<tr[^>]*>(.*?)</tr>", payload, re.S | re.I):
        cells = [
            re.sub(r"<[^>]+>", "", cell).replace("&nbsp;", " ").strip()
            for cell in re.findall(r"<td[^>]*>(.*?)</td>", row, re.S | re.I)
        ]

        if len(cells) <= index or not re.fullmatch(r"\d{2}/\d{2}/\d{4}", cells[0]):
            continue

        day, month, year = cells[0].split("/")
        stamp = f"{year}-{month}-{day}"

        try:
            rate = float(cells[index].replace(",", ""))
        except ValueError:
            continue

        if start <= stamp <= end:
            fixings[stamp] = rate / 100.0

    if not fixings:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    return fixings


async def _fetch_cnb(column: str, start: str, end: str) -> dict[str, float]:
    """Fetch a PRIBOR tenor from the CNB's yearly interbank fixing tables."""
    from datetime import datetime

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    index = int(column)
    last_year = int(end[:4])
    first_year = max(int(start[:4]), last_year - CNB_MAX_YEARS)
    fixings: dict[str, float] = {}

    async with aiohttp.ClientSession() as session:
        for year in range(first_year, last_year + 1):
            async with session.get(f"{CNB_PRIBOR_URL}?year={year}") as response:
                if response.status != 200:
                    continue

                payload = await response.text()

            for line in payload.splitlines():
                parts = line.split("|")

                if len(parts) <= index:
                    continue

                try:
                    day = (
                        datetime.strptime(parts[0].strip(), "%d %b %Y")
                        .date()
                        .isoformat()
                    )
                    rate = float(parts[index].strip())
                except ValueError:
                    continue

                if start <= day <= end:
                    fixings[day] = rate / 100.0

    if not fixings:
        raise OpenBBError(f"Unexpected fixings response from {CNB_PRIBOR_URL}.")

    return fixings


async def _fetch_bbk(series: str, start: str, end: str) -> dict[str, float]:
    """Fetch a Bundesbank daily money-market series from its CSV download."""
    import csv
    import io
    import re

    import aiohttp
    from openbb_core.app.model.abstract.error import OpenBBError

    url = f"{BBK_DOWNLOAD_URL}/{series}?format=csv&lang=en"

    async with (
        aiohttp.ClientSession() as session,
        session.get(url) as response,
    ):
        response.raise_for_status()
        text = await response.text()

    if not isinstance(text, str) or "," not in text:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    fixings: dict[str, float] = {}

    for row in csv.reader(io.StringIO(text)):
        if len(row) < 2 or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", row[0].strip()):
            continue

        value = row[1].strip()

        if not value:
            continue

        try:
            rate = float(value)
        except ValueError:
            continue

        day = row[0].strip()

        if start <= day <= end:
            fixings[day] = rate / 100.0

    if not fixings:
        raise OpenBBError(f"Unexpected fixings response from {url}.")

    return fixings


async def _fetch_rbnz(series: str, start: str, end: str) -> dict[str, float]:
    """Fetch NZIONA from the RBNZ B2 workbook via curl."""
    import asyncio
    import io

    import openpyxl
    from openbb_core.app.model.abstract.error import OpenBBError

    from openbb_cftc.utils.search import SEARCH_USER_AGENT

    process = await asyncio.create_subprocess_exec(
        "curl",
        "-sf",
        "--max-time",
        "30",
        "-A",
        SEARCH_USER_AGENT,
        RBNZ_B2_URL,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.PIPE,
    )  # noqa: S603, S607
    content, stderr = await process.communicate()

    if process.returncode != 0 or not content:
        raise OpenBBError(
            f"Unexpected fixings response from {RBNZ_B2_URL} -> "
            f"{stderr.decode(errors='ignore').strip()}"
        )

    try:
        book = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        sheet = book["Data"]
        column = next(
            (
                col
                for col in range(1, sheet.max_column + 1)
                if sheet.cell(row=5, column=col).value == series
            ),
            None,
        )
    except Exception as exc:
        raise OpenBBError(f"Unexpected fixings response from {RBNZ_B2_URL}.") from exc

    if column is None:
        raise OpenBBError(f"Unexpected fixings response from {RBNZ_B2_URL}.")

    fixings: dict[str, float] = {}

    for row_index in range(6, sheet.max_row + 1):
        day_cell = sheet.cell(row=row_index, column=1).value
        value_cell = sheet.cell(row=row_index, column=column).value

        if day_cell is None or value_cell is None:
            continue

        day = day_cell.date().isoformat()

        if start <= day <= end:
            fixings[day] = float(value_cell) / 100.0

    return fixings


async def _fetch_rba(series: str, start: str, end: str) -> dict[str, float]:
    """Fetch AONIA from the RBA's own F1 daily workbook, direct from rba.gov.au."""
    import io

    import aiohttp
    import openpyxl
    from openbb_core.app.model.abstract.error import OpenBBError

    async with (
        aiohttp.ClientSession() as session,
        session.get(RBA_F01D_URL) as response,
    ):
        response.raise_for_status()
        content = await response.read()

    try:
        book = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        sheet = book["Data"]
        column = next(
            (
                col
                for col in range(1, sheet.max_column + 1)
                if sheet.cell(row=11, column=col).value == series
            ),
            None,
        )
    except Exception as exc:
        raise OpenBBError(f"Unexpected fixings response from {RBA_F01D_URL}.") from exc

    if column is None:
        raise OpenBBError(f"Unexpected fixings response from {RBA_F01D_URL}.")

    fixings: dict[str, float] = {}

    for row_index in range(12, sheet.max_row + 1):
        day_cell = sheet.cell(row=row_index, column=1).value
        value_cell = sheet.cell(row=row_index, column=column).value

        if not isinstance(day_cell, dateType) or value_cell in (None, ""):
            continue

        day = day_cell.date().isoformat()

        if start <= day <= end:
            fixings[day] = float(value_cell) / 100.0

    return fixings


FULL_HISTORY_START = "2000-01-01"


async def get_fixings(
    index: str, start_date: dateType, end_date: dateType, use_cache: bool = True
) -> dict[dateType, float]:
    """Return an index's published fixings over a window, as decimals by effective date."""
    from datetime import datetime, timedelta, timezone

    from openbb_core.app.model.abstract.error import OpenBBError

    from openbb_cftc.utils import store

    source, path = FIXING_SOURCES[index]
    key = f"fixings-{index}"
    days = [
        (start_date + timedelta(days=offset)).isoformat()
        for offset in range((end_date - start_date).days + 1)
    ]
    today = datetime.now(timezone.utc).date().isoformat()
    cached = store.get_search_days(key, days) if use_cache else {}
    missing = [day for day in days if day not in cached]
    by_day: dict[str, list] = dict(cached)

    if missing:
        fetch_start = min(FULL_HISTORY_START, missing[0])
        fetch_end = today

        if source == "bcb":
            ten_years_back = (
                dateType.fromisoformat(today) - timedelta(days=3650)
            ).isoformat()
            fetch_start = max(fetch_start, ten_years_back)

        wide_days = [
            (dateType.fromisoformat(fetch_start) + timedelta(days=offset)).isoformat()
            for offset in range(
                (
                    dateType.fromisoformat(fetch_end)
                    - dateType.fromisoformat(fetch_start)
                ).days
                + 1
            )
        ]

        try:
            if source == "nyfed":
                rates = await _fetch_nyfed(path, fetch_start, fetch_end)
            elif source == "boc":
                rates = await _fetch_corra(fetch_start, fetch_end)
            elif source == "boe":
                rates = await _fetch_boe(path, fetch_start, fetch_end)
            elif source == "boj":
                rates = await _fetch_boj(path, fetch_start, fetch_end)
            elif source == "banxico":
                rates = await _fetch_banxico(path, fetch_start, fetch_end)
            elif source == "bcb":
                rates = await _fetch_bcb(path, fetch_start, fetch_end)
            elif source == "snb":
                rates = await _fetch_snb(path, fetch_start, fetch_end)
            elif source == "mas":
                rates = await _fetch_mas(path, fetch_start, fetch_end)
            elif source == "fbil":
                rates = await _fetch_fbil(path, fetch_start, fetch_end)
            elif source == "banrep":
                rates = await _fetch_banrep(path, fetch_start, fetch_end)
            elif source == "sarb":
                rates = await _fetch_sarb(path, fetch_start, fetch_end)
            elif source == "bcch":
                rates = await _fetch_bcch(path, fetch_start, fetch_end)
            elif source == "thaibma":
                rates = await _fetch_thaibma(path, fetch_start, fetch_end)
            elif source == "boi":
                rates = await _fetch_boi(path, fetch_start, fetch_end)
            elif source == "bbk":
                rates = await _fetch_bbk(path, fetch_start, fetch_end)
            elif source == "cnb":
                rates = await _fetch_cnb(path, fetch_start, fetch_end)
            elif source == "bnm":
                rates = await _fetch_bnm(path, fetch_start, fetch_end)
            elif source == "hkma":
                rates = await _fetch_hkma(path, fetch_start, fetch_end)
            elif source == "sarb_series":
                rates = await _fetch_sarb_series(path, fetch_start, fetch_end)
            elif source == "bok":
                rates = await _fetch_bok(path, fetch_start, fetch_end)
            elif source == "mnb":
                rates = await _fetch_mnb(path, fetch_start, fetch_end)
            elif source == "gpw":
                rates = await _fetch_gpw(path, fetch_start, fetch_end)
            elif source == "finwire":
                rates = await _fetch_finwire(path, fetch_start, fetch_end)
            elif source == "cfets":
                rates = await _fetch_cfets(path, fetch_start, fetch_end)
            elif source == "shibor":
                rates = await _fetch_shibor(path, fetch_start, fetch_end)
            elif source == "taibor":
                rates = await _fetch_taibor(path, fetch_start, fetch_end)
            elif source == "rbnz":
                rates = await _fetch_rbnz(path, fetch_start, fetch_end)
            elif source == "rba":
                rates = await _fetch_rba(path, fetch_start, fetch_end)
            else:
                rates = await _fetch_estr(fetch_start, fetch_end)
        except Exception as exc:  # noqa: BLE001
            if all(day >= today for day in missing):
                return {
                    dateType.fromisoformat(day): rows[0]["rate"]
                    for day, rows in by_day.items()
                    if rows
                }

            if isinstance(exc, OpenBBError):
                raise

            raise OpenBBError(f"Failed to fetch {index} fixings -> {exc}") from exc

        fetched: dict[str, list] = {day: [] for day in wide_days}

        for day, rate in rates.items():
            if day in fetched:
                fetched[day] = [{"rate": rate}]

        by_day.update({day: fetched[day] for day in missing})

        if use_cache:
            store.put_search_days(key, fetched)

    return {
        dateType.fromisoformat(day): rows[0]["rate"]
        for day, rows in by_day.items()
        if rows
    }


def compound_fixings(
    fixings: dict[dateType, float],
    start_date: dateType,
    end_date: dateType,
    basis: float = 360.0,
) -> float | None:
    """Compound published fixings over [start, end)."""
    if basis == 252.0:
        applicable = sorted(day for day in fixings if start_date <= day < end_date)

        if not applicable:
            return None

        factor = 1.0

        for day in applicable:
            factor *= 1.0 + fixings[day]

        return factor - 1.0

    dates = sorted(day for day in fixings if day < end_date)

    if not dates:
        return None

    applicable = [day for day in dates if day >= start_date]
    earlier = [day for day in dates if day < start_date]

    if not earlier and (not applicable or applicable[0] != start_date):
        return None

    if earlier:
        applicable = [earlier[-1], *applicable]

    factor = 1.0

    for index, day in enumerate(applicable):
        applied_from = max(day, start_date)
        applied_to = (
            min(applicable[index + 1], end_date)
            if index + 1 < len(applicable)
            else end_date
        )
        days = (applied_to - applied_from).days

        if days > 0:
            factor *= 1.0 + fixings[day] * days / basis

    return factor - 1.0
