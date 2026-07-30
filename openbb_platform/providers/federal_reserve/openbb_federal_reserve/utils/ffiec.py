"""FFIEC National Information Center (NIC) client."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from typing import Any

from openbb_federal_reserve.utils.curl_session import (
    get_session,
    reset_session as _reset,
)

BASE_URL = "https://www.ffiec.gov/npw"

QUARTER_END = {1: "0331", 2: "0630", 3: "0930", 4: "1231"}

_BHCPR_MONTH_TO_QUARTER = {
    "mar": 1,
    "march": 1,
    "jun": 2,
    "june": 2,
    "sep": 3,
    "sept": 3,
    "september": 3,
    "dec": 4,
    "december": 4,
}


def _warmup(session: Any) -> None:
    """Prime the Cloudflare cookie by hitting the NIC home page once."""
    session.get(f"{BASE_URL}/", timeout=30)


def _get_session() -> Any:
    """Return a warmed ``curl_cffi`` session that has the Cloudflare cookie."""
    return get_session("ffiec_nic", _warmup)


def reset_session() -> None:
    """Drop the cached session so the next call re-warms the Cloudflare cookie."""
    _reset("ffiec_nic")


def _fetch_bytes(path: str, referer: str | None = None) -> bytes:
    """Fetch a NIC URL as bytes, re-warming the session once on a 403."""
    url = f"{BASE_URL}/{path}"
    headers = {"Referer": referer or f"{BASE_URL}/"}
    session = _get_session()
    response = session.get(url, headers=headers, timeout=180)
    if response.status_code == 403:
        reset_session()
        response = _get_session().get(url, headers=headers, timeout=180)
    response.raise_for_status()
    return response.content


def _cached_bytes(
    key: str, path: str, referer: str | None = None, cadence: str = "daily"
) -> bytes:
    """Disk-cache a NIC download keyed by ``key``, expiring on ``cadence``."""
    from openbb_federal_reserve.utils.cache import cached, seconds_until_next_release

    return cached(
        ("ffiec", key),
        lambda: seconds_until_next_release(cadence),
        lambda: _fetch_bytes(path, referer),
    )


def _read_single_zip_member(content: bytes) -> bytes:
    """Return the bytes of the sole member of a ZIP archive."""
    archive = zipfile.ZipFile(io.BytesIO(content))
    name = archive.namelist()[0]
    with archive.open(name) as member:
        return member.read()


def available_bhcf_years() -> list[int]:
    """Return the FR Y-9 reporting years offered by the NIC, newest first."""
    import re

    html = _cached_bytes(
        "bhcf_years",
        "FinancialReport/FinancialDataDownload",
        f"{BASE_URL}/FinancialReport/FinancialDataDownload",
    ).decode("utf-8", "ignore")
    years = {int(y) for y in re.findall(r"<option[^>]*>\s*(\d{4})\s*</option>", html)}
    return sorted(years, reverse=True)


def fetch_bhcf(year: int, quarter: int) -> list[dict[str, str]]:
    """Download and parse an FR Y-9 BHCF quarterly file into per-institution rows.

    Parameters
    ----------
    year : int
        The four-digit reporting year.
    quarter : int
        The reporting quarter (1-4).

    Returns
    -------
    list[dict[str, str]]
        One caret-delimited record per institution, keyed by MDRM item code.
    """
    if quarter not in QUARTER_END:
        raise ValueError("quarter must be 1, 2, 3, or 4.")
    filename = f"BHCF{year}{QUARTER_END[quarter]}.ZIP"
    content = _cached_bytes(
        f"bhcf_{year}_{quarter}",
        f"FinancialReport/ReturnBHCFZipFiles?zipfilename={filename}",
        f"{BASE_URL}/FinancialReport/FinancialDataDownload",
    )
    text = _read_single_zip_member(content).decode("utf-8", "ignore")
    reader = csv.DictReader(io.StringIO(text), delimiter="^")
    return [row for row in reader if row.get("RSSD9001")]


_ATTRIBUTE_ENDPOINTS = {
    "active": "ReturnAttributesActiveZipFileCSV",
    "closed": "ReturnAttributesClosedZipFileCSV",
    "branches": "ReturnAttributesBranchesZipFileCSV",
}


def fetch_institutions(status: str = "active") -> list[dict[str, str]]:
    """Download and parse the NIC institution attributes CSV.

    Parameters
    ----------
    status : str
        One of ``"active"``, ``"closed"``, or ``"branches"``.

    Returns
    -------
    list[dict[str, str]]
        One record per institution keyed by NIC attribute name (``#ID_RSSD`` ...).
    """
    if status not in _ATTRIBUTE_ENDPOINTS:
        raise ValueError("status must be 'active', 'closed', or 'branches'.")
    content = _cached_bytes(
        f"attributes_{status}",
        f"FinancialReport/{_ATTRIBUTE_ENDPOINTS[status]}",
        f"{BASE_URL}/FinancialReport/DataDownload",
        cadence="quarterly" if status == "closed" else "daily",
    )
    text = _read_single_zip_member(content).decode("utf-8", "ignore")
    return list(csv.DictReader(io.StringIO(text)))


def entity_type(rssd_id: str) -> str | None:
    """Return an institution's NIC entity type (e.g. ``NAT``), best-effort."""
    from openbb_federal_reserve.utils.cache import cached, seconds_until_next_release

    def _producer() -> dict[str, str]:
        """Build a ``{rssd_id: entity_type}`` index from the NIC directory."""
        return {
            record["#ID_RSSD"]: record.get("ENTITY_TYPE", "")
            for record in fetch_institutions("active")
            if record.get("#ID_RSSD")
        }

    try:
        index = cached(
            "entity_types",
            lambda: seconds_until_next_release("daily"),
            _producer,
        )
    except Exception:  # noqa: BLE001
        return None
    return index.get(str(rssd_id)) or None


def rssd_names() -> dict[str, str]:
    """Return a cached ``{rssd_id: institution name}`` index."""
    from openbb_federal_reserve.utils.cache import cached, seconds_until_next_release

    def _producer() -> dict[str, str]:
        """Build the name index from the NIC active and closed directories."""
        index: dict[str, str] = {}
        for status in ("active", "closed"):
            try:
                records = fetch_institutions(status)
            except Exception:  # noqa: BLE001, S112
                continue
            for record in records:
                rssd = record.get("#ID_RSSD")
                name = (record.get("NM_SHORT") or record.get("NM_LGL") or "").strip()
                if rssd and name and rssd not in index:
                    index[rssd] = name
        return index

    try:
        return cached(
            "rssd_names",
            lambda: seconds_until_next_release("daily"),
            _producer,
        )
    except Exception:  # noqa: BLE001
        return {}


def fetch_relationships() -> list[dict[str, str]]:
    """Download and parse the NIC holding-company relationships CSV."""
    content = _cached_bytes(
        "relationships",
        "FinancialReport/ReturnRelationshipsZipFileCSV",
        f"{BASE_URL}/FinancialReport/DataDownload",
    )
    text = _read_single_zip_member(content).decode("utf-8", "ignore")
    return list(csv.DictReader(io.StringIO(text)))


def fetch_transformations() -> list[dict[str, str]]:
    """Download and parse the NIC institution transformations (mergers) CSV."""
    content = _cached_bytes(
        "transformations",
        "FinancialReport/ReturnTransformationZipFileCSV",
        f"{BASE_URL}/FinancialReport/DataDownload",
    )
    text = _read_single_zip_member(content).decode("utf-8", "ignore")
    return list(csv.DictReader(io.StringIO(text)))


def fetch_top_holders() -> list[dict[str, Any]]:
    """Return the ranked list of the largest holding companies by total assets."""
    import json

    content = _cached_bytes(
        "top_holders",
        "Institution/TopHolderList",
        f"{BASE_URL}/Institution/TopHoldings",
    )
    return json.loads(content.decode("utf-8", "ignore"))


def fetch_dictionary() -> dict[str, str]:
    """Return the MDRM item-code to short-description map for BHCF/structure items."""
    from openpyxl import load_workbook

    content = _cached_bytes(
        "financial_dictionary",
        "StaticData/DataDownload/Financial_Download_Dictionary.xlsx",
        f"{BASE_URL}/FinancialReport/FinancialDataDownload",
    )
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    descriptions: dict[str, str] = {}
    for sheet in workbook.sheetnames:
        for row in workbook[sheet].iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            code = str(row[0]).strip().upper()
            description = next(
                (str(cell).strip() for cell in row[1:] if isinstance(cell, str)), ""
            )
            if code and description:
                descriptions[code] = description
    workbook.close()
    return descriptions


def fetch_fry15_snapshots() -> list[dict[str, str]]:
    """List available FR Y-15 systemic-risk snapshot periods and their file URLs."""
    import re

    html = _cached_bytes(
        "fry15_index",
        "FinancialReport/FRY15Reports",
        f"{BASE_URL}/FinancialReport/FRY15Reports",
    ).decode("utf-8", "ignore")
    snapshots: list[dict[str, str]] = []
    hrefs = re.findall(r'href="(/npw/StaticData/Y15SnapShot/[^"]*All\.csv)"', html)
    for href in sorted(set(hrefs), reverse=True):
        match = re.search(r"/(\d{8})_(\d{8})_FRY15", href)
        if match:
            snapshots.append(
                {
                    "report_date": match.group(1),
                    "publish_date": match.group(2),
                    "url": href,
                }
            )
    return snapshots


def fetch_fry15(
    report_date: str | None = None, dataset: str = "indicators"
) -> list[dict[str, str]]:
    """Download and parse an FR Y-15 systemic-risk snapshot CSV.

    Parameters
    ----------
    report_date : str | None
        The ``YYYYMMDD`` report date; defaults to the most recent snapshot.
    dataset : str
        ``"indicators"`` for the published systemic-risk indicators (compact), or
        ``"all"`` for the complete set of FR Y-15 line items.

    Returns
    -------
    list[dict[str, str]]
        One record per institution keyed by MDRM item code. The CSVs are decoded
        with ``utf-8-sig`` so the indicators file's byte-order mark is stripped.
    """
    snapshots = fetch_fry15_snapshots()
    if not snapshots:
        return []
    selected = snapshots[0]
    if report_date:
        selected = next(
            (s for s in snapshots if s["report_date"] == report_date), selected
        )
    url = selected["url"]
    if dataset == "indicators":
        url = url.replace("All.csv", "Indicators.csv")
    content = _cached_bytes(
        f"fry15_{dataset}_{selected['report_date']}",
        url.removeprefix("/npw/"),
        f"{BASE_URL}/FinancialReport/FRY15Reports",
    )
    return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig", "ignore"))))


READY_REPORTS: dict[str, dict[str, str]] = {
    "FRY9C": {
        "name": "FR Y-9C Consolidated Financial Statements",
        "structure": "fry9c",
    },
    "FRY9LP": {
        "name": "FR Y-9LP Parent Company Only Financial Statements"
        " for Large Holding Companies",
        "structure": "fry9lp",
    },
    "FRY9SP": {
        "name": "FR Y-9SP Parent Company Only Financial Statements"
        " for Small Holding Companies",
        "structure": "fry9sp",
    },
    "FFIEC002": {
        "name": "FFIEC 002 Report of Assets and Liabilities of U.S. Branches"
        " and Agencies of Foreign Banks",
        "structure": "ffiec002",
    },
    "FFIEC101": {
        "name": "FFIEC 101 Regulatory Capital Reporting (Advanced Approaches)",
        "structure": "ffiec101",
    },
    "FFIEC102": {
        "name": "FFIEC 102 Market Risk Regulatory Report",
        "structure": "ffiec102",
    },
    "FRY15": {
        "name": "FR Y-15 Banking Organization Systemic Risk Report",
        "structure": "fry15",
    },
    "FRQ1": {
        "name": "FR Q-1 Capital and Asset Report for Supervised Insurance"
        " Organizations (Building Block Approach)",
        "structure": "frq1",
    },
}

_REPORT_METADATA_ROWS = {
    "Institution Name": "institution_name",
    "Report Date": "report_date",
    "ID_RSSD": "rssd_id",
}


def fetch_financial_report(report_code: str, rssd_id: str, date: str) -> dict[str, Any]:
    """Fetch one institution's regulatory report as an MDRM-keyed fact map.

    Downloads the per-institution ``ReturnFinancialReportCSV`` payload (the FFIEC
    NIC "ItemName,Description,Value" listing) for the given report code, RSSD
    identifier, and quarter-end date, and parses it into the institution's name,
    report date, and a ``{mdrm: value}`` map of every value-bearing line item.
    Cached daily. An institution that did not file the report for the period
    returns an empty ``facts`` map (the CSV carries only the metadata rows).

    Parameters
    ----------
    report_code : str
        The ``rpt`` code understood by the endpoint (for example ``"FRY9C"``).
    rssd_id : str
        The reporting entity's RSSD identifier.
    date : str
        The reporting period as ``YYYYMMDD`` (a calendar quarter-end).

    Returns
    -------
    dict[str, Any]
        ``{"institution_name", "report_date", "rssd_id", "facts", "descriptions"}``
        where ``facts`` is the ``{mdrm: value}`` map of value-bearing line items
        and ``descriptions`` is the ``{mdrm: csv description}`` map of their filed
        captions, used to label codes a committed structure does not enumerate.
    """
    from openbb_federal_reserve.utils.cache import cached, seconds_until_next_release

    code = str(report_code).strip().upper()
    rssd = str(rssd_id).strip()
    period = str(date).strip()

    def _producer() -> dict[str, Any]:
        """Fetch and parse the per-institution CSV into a fact map."""
        raw = _fetch_bytes(
            "FinancialReport/ReturnFinancialReportCSV"
            f"?rpt={code}&id={rssd}&dt={period}",
            referer=f"{BASE_URL}/FinancialReport/FinancialDataDownload",
        )
        return _parse_financial_report_csv(raw.decode("utf-8", "replace"))

    return cached(
        ("financial_report", code, rssd, period),
        lambda: seconds_until_next_release("daily"),
        _producer,
    )


_PROFILE_REPORT_BLOCK = re.compile(
    r"if\(\s*(?P<count>\d+)\s*>\s*0\s*\)\s*\{\s*"
    r"var\s+dateCards\s*=\s*(?P<cards>\[.*?\])\s*;\s*"
    r"var\s+idRssd\s*=\s*\d+\s*;\s*"
    r"buildOptionCards\(\s*dateCards\s*,\s*['\"](?P<code>[A-Z0-9]+)['\"]",
    re.DOTALL,
)

_MONTH_DAY_TO_QUARTER = {"3/31": 1, "6/30": 2, "9/30": 3, "12/31": 4}

_PROFILE_REPORT_TITLE = re.compile(
    r'<button[^>]*href="#[A-Z0-9]+"[^>]*>\s*'
    r"(?P<title>[^<]*?\((?P<paren>[A-Z][^)]*)\)[^<]*?)\s*<",
    re.DOTALL,
)


def _parse_profile_report_names(html: str) -> dict[str, str]:
    """Map each report code to its official NIC name from the profile page."""
    names: dict[str, str] = {}
    for match in _PROFILE_REPORT_TITLE.finditer(html):
        code = re.sub(r"[^A-Z0-9]", "", match.group("paren").upper())
        title = re.sub(r"\s+", " ", match.group("title")).strip()
        if code and code not in names:
            names[code] = title
    return names


def fetch_institution_financial_reports(
    rssd_id: str,
) -> dict[str, dict[str, Any]]:
    """Return the financial reports an institution files, with names and periods.

    Parses the NIC Institution Profile page, whose embedded ``buildOptionCards``
    blocks list, for each filed report series, every period the institution has
    on file, and whose collapsible toggles carry each series' official NIC name.
    Reports the firm does not file (an empty or zero-guarded block) are omitted.
    Cached daily.

    Parameters
    ----------
    rssd_id : str
        The institution's RSSD identifier.

    Returns
    -------
    dict[str, dict[str, Any]]
        A ``{report_code: {"name", "periods"}}`` map where ``name`` is the report's
        official NIC name (or ``None`` when the profile lacks it) and each period is
        ``{"year", "quarter", "month_day", "pdf_available"}``, newest first.
    """
    import json

    from openbb_federal_reserve.utils.cache import cached, seconds_until_next_release

    rssd = str(rssd_id).strip()

    def _producer() -> dict[str, dict[str, Any]]:
        """Fetch the profile page and parse its filed-report blocks and names."""
        html = _fetch_bytes(
            f"Institution/Profile/{rssd}",
            referer=f"{BASE_URL}/Institution/Profile/{rssd}",
        ).decode("utf-8", "ignore")
        official_names = _parse_profile_report_names(html)
        reports: dict[str, dict[str, Any]] = {}
        for match in _PROFILE_REPORT_BLOCK.finditer(html):
            if int(match.group("count")) <= 0:
                continue
            try:
                date_cards = json.loads(match.group("cards"))
            except json.JSONDecodeError:  # pragma: no cover
                continue
            periods: list[dict[str, Any]] = []
            for card in date_cards:
                year = int(card["year"])
                for option in card.get("dateOptions", []):
                    month_day = option["monthDay"]
                    quarter = _MONTH_DAY_TO_QUARTER.get(month_day)
                    if quarter is None:
                        continue
                    periods.append(
                        {
                            "year": year,
                            "quarter": quarter,
                            "month_day": month_day,
                            "pdf_available": bool(option.get("pdfAvailable")),
                        }
                    )
            if periods:
                periods.sort(key=lambda p: (p["year"], p["quarter"]), reverse=True)
                code = match.group("code")
                reports[code] = {
                    "name": official_names.get(code),
                    "periods": periods,
                }
        return reports

    return cached(
        ("institution_financial_reports", rssd),
        lambda: seconds_until_next_release("daily"),
        _producer,
    )


def _normalize_report_date(value: str) -> str | None:
    """Normalize a CSV report-date cell to ``YYYYMMDD``."""
    from datetime import datetime

    text = (value or "").strip()
    if re.fullmatch(r"\d{8}", text):
        return text
    try:
        return datetime.strptime(text, "%B %d, %Y").strftime("%Y%m%d")
    except ValueError:
        return None


def _parse_financial_report_csv(text: str) -> dict[str, Any]:
    """Parse a ``ReturnFinancialReportCSV`` payload into identity plus facts."""
    rows = list(csv.reader(io.StringIO(text)))
    if not rows or rows[0][:3] != ["ItemName", "Description", "Value"]:
        return {
            "institution_name": None,
            "report_date": None,
            "rssd_id": None,
            "facts": {},
            "descriptions": {},
        }
    identity: dict[str, str | None] = {
        "institution_name": None,
        "report_date": None,
        "rssd_id": None,
    }
    facts: dict[str, str] = {}
    descriptions: dict[str, str] = {}
    for row in rows[1:]:
        if len(row) < 3:
            continue
        name = row[0].strip()
        value = row[2].strip()
        key = _REPORT_METADATA_ROWS.get(name)
        if key is not None:
            # non-empty value wins.
            if key == "report_date":
                normalized = _normalize_report_date(value)
                if normalized and (
                    identity[key] is None or re.fullmatch(r"\d{8}", normalized)
                ):
                    identity[key] = normalized
            elif identity[key] is None:
                identity[key] = value or None
            continue
        if re.fullmatch(r"[A-Z]{4}[A-Z0-9]{4}", name):
            facts[name] = value
            descriptions[name] = row[1].strip()
    return {**identity, "facts": facts, "descriptions": descriptions}


def _parse_bhcpr_name(name: str) -> dict[str, int] | None:
    """Parse a BHCPR PDF filename into peer group, year, and quarter."""
    import re

    legacy = re.match(r"PeerGroup_(\d+)_([A-Za-z]+)(\d{4})\.pdf$", name, re.IGNORECASE)
    if legacy:
        quarter = _BHCPR_MONTH_TO_QUARTER.get(legacy.group(2).lower())
        if quarter is None:
            return None
        return {
            "peer_group": int(legacy.group(1)),
            "year": int(legacy.group(3)),
            "quarter": quarter,
        }

    current = re.match(
        r"BHCPR_PeerGrp(\d+)_(\d{4})(\d{2})\d{2}\.pdf$", name, re.IGNORECASE
    )
    if current:
        quarter = {3: 1, 6: 2, 9: 3, 12: 4}.get(int(current.group(3)))
        if quarter is None:
            return None
        return {
            "peer_group": int(current.group(1)),
            "year": int(current.group(2)),
            "quarter": quarter,
        }

    return None


def list_bhcpr_reports() -> list[dict[str, Any]]:
    """List BHCPR peer-group report PDFs with parsed peer group, year, and quarter."""
    import re

    html = _cached_bytes(
        "bhcpr_index",
        "FinancialReport/BHCPRReports",
        f"{BASE_URL}/FinancialReport/BHCPRReports",
    ).decode("utf-8", "ignore")
    hrefs = re.findall(r'href="(/npw/StaticData/bhcpRRPT/[^"]+\.pdf)"', html)
    reports: list[dict[str, Any]] = []
    for href in sorted(set(hrefs)):
        name = href.rsplit("/", 1)[-1]
        parsed = _parse_bhcpr_name(name)
        if parsed is None:
            continue
        month_day = QUARTER_END[parsed["quarter"]]
        reports.append(
            {
                "name": name,
                "url": f"https://www.ffiec.gov{href}",
                "peer_group": parsed["peer_group"],
                "year": parsed["year"],
                "quarter": parsed["quarter"],
                "period_end": f"{parsed['year']}-{month_day[:2]}-{month_day[2:]}",
            }
        )
    return reports


def download_bhcpr_pdf(url: str) -> bytes:
    """Download a BHCPR peer-group report PDF through the warmed FFIEC session."""
    from urllib.parse import urlparse

    from openbb_core.provider.utils.errors import OpenBBError

    parsed = urlparse(url)
    if (
        parsed.scheme != "https"
        or parsed.hostname not in {"www.ffiec.gov", "ffiec.gov"}
        or not parsed.path.startswith("/npw/StaticData/bhcpRRPT/")
        or not parsed.path.lower().endswith(".pdf")
    ):
        raise OpenBBError(f"Invalid BHCPR report URL provided for download -> {url}")
    return _fetch_bytes(
        parsed.path[len("/npw/") :],
        referer=f"{BASE_URL}/FinancialReport/BHCPRReports",
    )


def download_financial_report_pdf(url: str) -> bytes:
    """Download a filed FFIEC financial-report PDF through the warmed session."""
    from urllib.parse import urlparse

    from openbb_core.provider.utils.errors import OpenBBError

    parsed = urlparse(url)
    if (
        parsed.scheme != "https"
        or parsed.hostname not in {"www.ffiec.gov", "ffiec.gov"}
        or parsed.path != "/npw/FinancialReport/ReturnFinancialReportPDF"
    ):
        raise OpenBBError(
            f"Invalid financial report PDF URL provided for download -> {url}"
        )
    path = parsed.path[len("/npw/") :]
    if parsed.query:
        path = f"{path}?{parsed.query}"
    return _fetch_bytes(
        path,
        referer=f"{BASE_URL}/FinancialReport/FinancialDataDownload",
    )


def fetch_bhcpr(rssd_id: str, date: str) -> dict[str, Any]:
    """Fetch one holding company's coded BHCPR CSV for a reporting period.

    Downloads the per-institution ``ReturnFinancialReportCSV?rpt=BHCPR`` payload
    for the RSSD identifier and quarter-end date and parses it into the
    institution's identity, the reporting-period dates, and every value keyed by
    its MDRM code. The committed schema binds each report line item to its code,
    so this coded CSV -- not the PDF -- is the value source. Cached daily. A
    holding company that did not file the BHCPR returns an empty ``values`` map.

    Parameters
    ----------
    rssd_id : str
        The reporting holding company's RSSD identifier.
    date : str
        The reporting period as ``YYYYMMDD`` (a calendar quarter-end).

    Returns
    -------
    dict[str, Any]
        ``{"identity", "periods", "values", "descriptions"}`` from
        ``parse_bhcpr_csv``: ``values`` maps each base code (e.g. ``BHSR028``) to
        its ``{period-suffix: value}`` series.
    """
    from openbb_federal_reserve.utils.bhcpr_csv import parse_bhcpr_csv
    from openbb_federal_reserve.utils.cache import cached, seconds_until_next_release

    rssd = str(rssd_id).strip()
    period = str(date).strip()

    def _producer() -> dict[str, Any]:
        """Download the per-institution BHCPR CSV and parse it into coded values."""
        raw = _fetch_bytes(
            f"FinancialReport/ReturnFinancialReportCSV?rpt=BHCPR&id={rssd}&dt={period}",
            referer=f"{BASE_URL}/FinancialReport/FinancialDataDownload",
        )
        return parse_bhcpr_csv(raw)

    return cached(
        ("bhcpr_data", rssd, period),
        lambda: seconds_until_next_release("daily"),
        _producer,
    )


def _top_tier_holder(rssd: str) -> str | None:
    """Return the RSSD of an institution's top-tier holder from the NIC hierarchy."""
    from openbb_federal_reserve.utils.cache import cached, seconds_until_next_release

    def _producer() -> str | None:
        """Post to the NIC BuildTier hierarchy and read the top-tier RSSD."""
        try:
            response = _get_session().post(
                f"{BASE_URL}/Institution/BuildTier",
                data={
                    "RssdID": rssd,
                    "ProfileDtStart": "1/1/1900 12:00:00 AM",
                    "ProfileDtEnd": "12/31/9999 12:00:00 AM",
                },
                headers={"Referer": f"{BASE_URL}/Institution/Profile/{rssd}"},
                timeout=60,
            )
            top = response.json().get("top_tier_id_rssd")
        except Exception:  # noqa: BLE001
            return None
        return str(top) if top else None

    return cached(
        ("bhcpr_top_holder", rssd),
        lambda: seconds_until_next_release("daily"),
        _producer,
    )


def resolve_bhcpr_holder(rssd_id: str) -> str:
    """Resolve an RSSD to the top-tier holding company that files the BHCPR."""
    rssd = str(rssd_id).strip()
    if not rssd:
        return rssd
    try:
        reports = fetch_institution_financial_reports(rssd)
    except Exception:  # noqa: BLE001
        reports = None
    if isinstance(reports, dict) and "BHCPR" in reports:
        return rssd
    return _top_tier_holder(rssd) or rssd
