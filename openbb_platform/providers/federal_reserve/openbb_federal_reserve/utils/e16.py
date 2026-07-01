"""FFIEC E.16 Country Exposure Lending Survey (FFIEC 009) client and parser.

The E.16 statistical release publishes aggregate country-exposure tables for
U.S. banks as a cleansed Excel workbook, zipped, at a static URL keyed by the
release quarter. Each sheet is one table for one institution group (All Banks,
Large Financial Institutions, All Others); rows are countries grouped by region,
columns are exposure measures under a merged multi-level header.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import (
    date as dateType,
    datetime,
)
from typing import Any

from openbb_federal_reserve.utils.curl_session import get_session

BASE_URL = "https://www.ffiec.gov/sites/default/files/data/e16"
PAGE_URL = "https://www.ffiec.gov/data/e16"

# Friendly group key -> the workbook's sheet-name prefix.
GROUPS = {"all_banks": "All Banks", "lfi": "LFI", "all_others": "All Others"}
TABLES = ("1", "2", "3", "4.1", "4.2")


def _get_session() -> Any:
    """Return a browser-impersonating ``curl_cffi`` session."""
    return get_session("e16")


def _file_quarter(report_date: dateType) -> str:
    """Return the data file's ``YYYYMM`` - the report quarter end plus 3 months."""
    month = report_date.month + 3
    year = report_date.year + (month - 1) // 12
    return f"{year}{((month - 1) % 12) + 1:02d}"


def fetch_e16(report_date: dateType | None = None) -> bytes:
    """Download the E.16 data ZIP for a report date (latest if None), cached."""
    from openbb_federal_reserve.utils.cache import cached, seconds_until_next_release

    def _producer() -> bytes:
        """Resolve the quarter's data-file URL and download it."""
        session = _get_session()
        if report_date:
            quarter = _file_quarter(report_date)
        else:
            page = session.get(PAGE_URL, timeout=60).text
            quarters = re.findall(r"/E16_(\d{6})\.zip", page)
            if not quarters:
                return b""
            quarter = max(quarters)
        return session.get(f"{BASE_URL}/E16_{quarter}.zip", timeout=180).content

    return cached(
        ("e16", report_date.isoformat() if report_date else "latest"),
        lambda: seconds_until_next_release("quarterly"),
        _producer,
    )


def _merged_values(worksheet: Any) -> dict[tuple[int, int], Any]:
    """Fill every cell of a merged range with the range's top-left value."""
    values: dict[tuple[int, int], Any] = {}
    for cells in worksheet.merged_cells.ranges:
        value = worksheet.cell(cells.min_row, cells.min_col).value
        for row in range(cells.min_row, cells.max_row + 1):
            for col in range(cells.min_col, cells.max_col + 1):
                values[(row, col)] = value
    return values


def _is_number(value: Any) -> bool:
    """Return whether a cell value parses as a number."""
    try:
        float(str(value).replace(",", "").strip())
        return True
    except (TypeError, ValueError):
        return False


def _period(workbook: Any) -> dateType | None:
    """Read the report period from the workbook's cover sheet."""
    for row in workbook[workbook.sheetnames[0]].iter_rows(max_row=12, values_only=True):
        for cell in row:
            match = re.search(r"Period:\s*([A-Za-z]+ \d+, \d{4})", str(cell or ""))
            if match:
                try:
                    return datetime.strptime(match.group(1), "%B %d, %Y").date()
                except ValueError:
                    return None
    return None


def parse_sheet(
    content: bytes, sheet: str
) -> tuple[list[dict[str, Any]], dateType | None]:
    """Parse one E.16 table sheet into long-format country-exposure rows.

    Returns ``(rows, period)`` where each row is ``{country_group, country,
    label, value}``. Labels combine the merged multi-level column header (the
    table-wide title row is dropped, footnote markers stripped); countries are
    grouped under their region header. The survey reports amounts in millions of
    dollars; values are scaled to dollars.
    """
    from openpyxl import load_workbook

    archive = zipfile.ZipFile(io.BytesIO(content))
    member = next((m for m in archive.namelist() if m.lower().endswith(".xlsx")), None)
    if member is None:
        return [], None
    workbook = load_workbook(io.BytesIO(archive.read(member)), data_only=True)
    period = _period(workbook)
    if sheet not in workbook.sheetnames:
        return [], period
    worksheet = workbook[sheet]
    merged = _merged_values(worksheet)

    def cell(row: int, col: int) -> Any:
        """Return a cell's value, resolving merged ranges."""
        value = merged.get((row, col))
        return value if value is not None else worksheet.cell(row, col).value

    columns = range(3, worksheet.max_column + 1)
    data_rows = [
        row
        for row in range(1, worksheet.max_row + 1)
        if cell(row, 2) and any(_is_number(cell(row, col)) for col in columns)
    ]
    if not data_rows:
        return [], period

    # Header rows are the text rows above the data, excluding the code row.
    header_rows = []
    for row in range(1, data_rows[0]):
        labels = [
            str(cell(row, col)).strip()
            for col in columns
            if cell(row, col) and str(cell(row, col)).strip() not in ("", "\xa0")
        ]
        if labels and not all(re.fullmatch(r"\(.*\)", label) for label in labels):
            header_rows.append(row)
    leaf = header_rows[-1]
    value_cols = [
        col
        for col in columns
        if cell(leaf, col) and str(cell(leaf, col)).strip() not in ("", "\xa0")
    ]

    def label(col: int) -> str:
        """Build a column's full label from its merged header hierarchy."""
        parts: list[str] = []
        for row in header_rows:
            # A value spanning every column is the table title, not a header.
            if len({str(cell(row, other)) for other in value_cols}) <= 1:
                continue
            value = cell(row, col)
            if value and str(value).strip() not in ("", "\xa0"):
                clean = re.sub(r"\s*/\d+", "", re.sub(r"\s+", " ", str(value))).strip()
                if clean and clean not in parts:
                    parts.append(clean)
        return " - ".join(parts)

    labels = {col: label(col) for col in value_cols}
    rows: list[dict[str, Any]] = []
    group: str | None = None
    for row in range(leaf + 1, worksheet.max_row + 1):
        name = cell(row, 2)
        if not name:
            continue
        name = str(name).strip()
        values = [
            (col, cell(row, col)) for col in value_cols if _is_number(cell(row, col))
        ]
        if not values:
            group = name
            continue
        for col, value in values:
            rows.append(
                {
                    "country_group": group,
                    "country": name,
                    "label": labels[col],
                    "value": float(str(value).replace(",", "").strip()) * 1_000_000,
                }
            )
    return rows, period
