"""Tests for the FFIEC E.16 Country Exposure fetcher model."""

import io
import zipfile
from datetime import date

import pytest
from openbb_core.provider.utils.errors import EmptyDataError
from openpyxl import Workbook

from openbb_federal_reserve.models.country_exposure import (
    FederalReserveCountryExposureData,
    FederalReserveCountryExposureFetcher,
    FederalReserveCountryExposureQueryParams,
)
from openbb_federal_reserve.utils.e16 import GROUPS, TABLES


def _zip() -> bytes:
    """Build an E.16 data ZIP: a cover sheet and one table sheet."""
    workbook = Workbook()
    cover = workbook.active
    cover.title = "E16_009"
    cover["A1"] = "Statistical Release E.16  Period: December 31, 2024"

    sheet = workbook.create_sheet("All Banks - Table 1")
    sheet.cell(1, 2, "Country Exposure Survey Table 1")
    sheet.merge_cells("B1:F1")  # the table-wide title, dropped from labels
    sheet.cell(2, 3, "Claims")
    sheet.merge_cells("C2:D2")
    sheet.cell(2, 5, "Off-Balance Sheet")
    sheet.merge_cells("E2:F2")
    sheet.cell(3, 3, "Cross-border Claims /2")
    sheet.cell(3, 4, "Total")
    sheet.cell(3, 5, "Unused Commitments")
    sheet.cell(3, 6, "Guarantees")
    sheet.cell(4, 2, "G-10 and Luxembourg")  # region header (no values)
    for row, country, values in (
        (5, "CANADA", (179188, 82747, 78777, 28438)),
        (6, "FRANCE", (152588, 20399, 220289, 6169)),
    ):
        sheet.cell(row, 2, country)
        for offset, value in enumerate(values):
            sheet.cell(row, 3 + offset, value)

    book = io.BytesIO()
    workbook.save(book)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("Dec 31 2024 - E16 (009)_Cleansed.xlsx", book.getvalue())
    return buffer.getvalue()


def _patch(monkeypatch):
    """Point the E.16 download at the fixture workbook."""
    monkeypatch.setattr(
        "openbb_federal_reserve.utils.e16.fetch_e16", lambda report_date=None: _zip()
    )


class TestExtractData:
    """Tests for the workbook download and parse."""

    def test_parses_table(self, monkeypatch):
        """Countries pivot to wide rows with region grouping and measure columns."""
        _patch(monkeypatch)
        query = FederalReserveCountryExposureFetcher.transform_query(
            {"group": "all_banks", "table": "1"}
        )
        rows = FederalReserveCountryExposureFetcher.extract_data(query, None)
        result = FederalReserveCountryExposureFetcher.transform_data(query, rows)
        assert all(isinstance(r, FederalReserveCountryExposureData) for r in result)
        assert all(r.date == date(2024, 12, 31) for r in result)
        assert all(r.country_group == "G-10 and Luxembourg" for r in result)
        by_country = {r.country: r.model_dump() for r in result}
        # One row per country; section + leaf combine into measure columns.
        assert (
            by_country["CANADA"]["Claims - Cross-border Claims"] == 179188 * 1_000_000
        )
        assert (
            by_country["CANADA"]["Off-Balance Sheet - Unused Commitments"]
            == 78777 * 1_000_000
        )
        assert by_country["FRANCE"]["Claims - Total"] == 20399 * 1_000_000

    def test_country_filter(self, monkeypatch):
        """The country filter keeps only the matching country."""
        _patch(monkeypatch)
        query = FederalReserveCountryExposureFetcher.transform_query(
            {"country": "france"}
        )
        rows = FederalReserveCountryExposureFetcher.extract_data(query, None)
        assert {row["country"] for row in rows} == {"FRANCE"}

    def test_country_filter_with_no_match_raises(self, monkeypatch):
        """A country filter that matches nothing raises ``EmptyDataError``."""
        _patch(monkeypatch)
        query = FederalReserveCountryExposureFetcher.transform_query(
            {"country": "atlantis"}
        )
        with pytest.raises(EmptyDataError):
            FederalReserveCountryExposureFetcher.extract_data(query, None)

    def test_empty_response_raises(self, monkeypatch):
        """A non-ZIP response raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_federal_reserve.utils.e16.fetch_e16",
            lambda report_date=None: b"",
        )
        query = FederalReserveCountryExposureFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveCountryExposureFetcher.extract_data(query, None)


class TestWidgetOptions:
    """Tests for the human-readable ``group`` and ``table`` dropdown labels."""

    def test_group_options_have_human_labels(self):
        """Every reporting-group code maps to a human label, not a bare code."""
        options = FederalReserveCountryExposureQueryParams.__json_schema_extra__[
            "group"
        ]["x-widget_config"]["options"]
        assert {o["value"] for o in options} == set(GROUPS)
        for option in options:
            assert option["label"] and option["label"] != option["value"]

    def test_table_options_have_human_labels(self):
        """Every survey-table code maps to a human label, not a bare code."""
        options = FederalReserveCountryExposureQueryParams.__json_schema_extra__[
            "table"
        ]["x-widget_config"]["options"]
        assert {o["value"] for o in options} == set(TABLES)
        for option in options:
            assert option["label"] and option["label"] != option["value"]
