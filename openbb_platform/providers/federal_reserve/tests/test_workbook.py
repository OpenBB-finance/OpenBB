"""Tests for the shared workbook utilities."""

from datetime import date, datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from pandas import DataFrame

from openbb_federal_reserve.utils.workbook import (
    melt_sheet,
    pivot_wide,
    round_value,
    sanitize_xlsx_core_dates,
)

_CORE_PATH = "docProps/core.xml"

# A core.xml whose timestamps use a space-padded hour (``T 2:39:25``), as written
# by SAS-generated Philadelphia Fed workbooks. openpyxl cannot parse the hour and
# falls back to a date-only value that then fails its datetime descriptor.
_MALFORMED_CORE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    "<cp:coreProperties"
    ' xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"'
    ' xmlns:dc="http://purl.org/dc/elements/1.1/"'
    ' xmlns:dcterms="http://purl.org/dc/terms/"'
    ' xmlns:dcmitype="http://purl.org/dc/dcmitype/"'
    ' xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
    '<dcterms:created xsi:type="dcterms:W3CDTF">2026-05-13T 2:39:25-04:00</dcterms:created>\n'
    '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-05-13T 2:39:25-04:00</dcterms:modified>\n'
    "</cp:coreProperties>"
)


def _xlsx_bytes() -> bytes:
    """Return well-formed single-sheet ``.xlsx`` bytes with a standard core.xml."""
    buffer = BytesIO()
    DataFrame({"value": [1]}).to_excel(buffer, index=False)
    return buffer.getvalue()


def _rewrite_core(content: bytes, core_xml: str | None) -> bytes:
    """Return ``content`` with docProps/core.xml replaced, or removed when ``None``."""
    out = BytesIO()
    with ZipFile(BytesIO(content)) as zin, ZipFile(out, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == _CORE_PATH:
                if core_xml is None:
                    continue
                zout.writestr(item, core_xml.encode("utf-8"))
            else:
                zout.writestr(item, zin.read(item.filename))
    return out.getvalue()


class TestSanitizeXlsxCoreDates:
    """Tests for ``sanitize_xlsx_core_dates``."""

    def test_missing_core_properties_returns_input(self):
        """A workbook without docProps/core.xml is returned unchanged."""
        content = _rewrite_core(_xlsx_bytes(), None)
        assert sanitize_xlsx_core_dates(content) is content

    def test_well_formed_workbook_returned_unchanged(self):
        """A workbook with valid timestamps is returned as the same object."""
        content = _xlsx_bytes()
        assert sanitize_xlsx_core_dates(content) is content

    def test_space_padded_hour_is_repaired(self):
        """A space-padded hour is zero-padded so openpyxl can load the workbook."""
        malformed = _rewrite_core(_xlsx_bytes(), _MALFORMED_CORE)

        with pytest.raises(TypeError):
            load_workbook(BytesIO(malformed), read_only=True)

        repaired = sanitize_xlsx_core_dates(malformed)
        assert repaired != malformed

        workbook = load_workbook(BytesIO(repaired), read_only=True)
        assert workbook.properties.created == datetime(2026, 5, 13, 2, 39, 25)
        workbook.close()


def _save_sheet(rows: list[list]) -> bytes:
    """Build single-sheet ``.xlsx`` bytes from a list of rows."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class TestRoundValue:
    """Tests for ``round_value``."""

    def test_rounds_to_two_decimals_by_default(self):
        """A noisy float rounds to two decimals at the default precision."""
        assert round_value(2.7152000003) == 2.72

    def test_small_value_widens_to_keep_significant_digits(self):
        """A value that would round to zero widens to its first significant digits."""
        assert round_value(0.0030034608406641) == 0.003
        assert round_value(0.0019010331209758) == 0.0019

    def test_zero_returns_zero_float(self):
        """An exact zero short-circuits to ``0.0`` without widening."""
        assert round_value(0) == 0.0

    def test_non_finite_passes_through(self):
        """A non-finite float is returned unchanged."""
        assert round_value(float("nan")) != round_value(float("nan"))
        assert round_value(float("inf")) == float("inf")

    def test_non_numeric_and_bool_pass_through(self):
        """Strings, ``None``, and booleans are returned unchanged."""
        assert round_value("United States") == "United States"
        assert round_value(None) is None
        assert round_value(True) is True


class TestPivotWide:
    """Tests for ``pivot_wide``."""

    def test_one_row_per_index_with_series_columns(self):
        """Each distinct index becomes one row carrying every series column."""
        records = [
            {"date": date(2024, 1, 31), "series": "a", "value": 1.0},
            {"date": date(2024, 1, 31), "series": "b", "value": 2.0},
            {"date": date(2024, 2, 29), "series": "a", "value": 3.0},
        ]
        rows = pivot_wide(records)
        assert rows == [
            {"date": date(2024, 1, 31), "a": 1.0, "b": 2.0},
            {"date": date(2024, 2, 29), "a": 3.0},
        ]

    def test_all_none_row_dropped_partial_row_kept(self):
        """A row whose every series value is ``None`` is dropped; a partial row stays."""
        records = [
            {"date": date(1997, 1, 1), "series": "a", "value": None},
            {"date": date(1997, 1, 1), "series": "b", "value": None},
            {"date": date(1997, 2, 1), "series": "a", "value": 5.0},
            {"date": date(1997, 2, 1), "series": "b", "value": None},
        ]
        rows = pivot_wide(records)
        assert rows == [{"date": date(1997, 2, 1), "a": 5.0, "b": None}]

    def test_tuple_index_keeps_index_fields_out_of_emptiness(self):
        """With a tuple index, only non-index values count toward an all-None drop."""
        records = [
            {"date": date(2024, 1, 31), "country": "US", "series": "a", "value": None},
            {"date": date(2024, 1, 31), "country": "JP", "series": "a", "value": 7.0},
        ]
        rows = pivot_wide(records, index=("date", "country"))
        assert rows == [
            {"date": date(2024, 1, 31), "country": "JP", "a": 7.0},
        ]


class TestMeltSheet:
    """Tests for ``melt_sheet``."""

    def test_no_date_column_returns_empty(self):
        """A sheet whose first column never resolves to a date yields no records."""
        content = _save_sheet([["Note", "Value"], ["foo", 1], ["bar", 2]])
        assert melt_sheet(content, "Sheet1") == []

    def test_grouped_header_and_labels_remap(self):
        """Group headers prefix measures and the labels map renames a series."""
        content = _save_sheet(
            [
                ["Source: example"],
                [None, "Region A", "Region A", "Note: ignore"],
                ["Date", "Corn", "Wheat", "Last updated 2025"],
                [datetime(2024, 1, 31), 1.0, 2.0, "x"],
                ["2024-02-29", 3.0, None, "y"],
            ]
        )
        records = melt_sheet(
            content,
            "Sheet1",
            group_header=True,
            labels={"Region A - Corn": "Corn (Region A)"},
        )
        by_key = {(r["date"], r["series"]): r["value"] for r in records}
        assert by_key[(date(2024, 1, 31), "Corn (Region A)")] == 1.0
        assert by_key[(date(2024, 1, 31), "Region A - Wheat")] == 2.0
        assert by_key[(date(2024, 2, 29), "Region A - Wheat")] is None

    def test_key_columns_carried_as_fields(self):
        """A key column is carried into each record instead of melted as a series."""
        content = _save_sheet(
            [
                ["Date", "Vintage", "Value"],
                [datetime(2024, 1, 31), "first", 10.0],
                [datetime(2024, 1, 31), None, 11.0],
            ]
        )
        records = melt_sheet(content, "Sheet1", key_columns=("Vintage",))
        assert {r["series"] for r in records} == {"Value"}
        by_value = {r["value"]: r["vintage"] for r in records}
        assert by_value[10.0] == "first"
        assert by_value[11.0] is None

    def test_string_dates_are_detected(self):
        """A leading column of date strings is recognized as the data start."""
        content = _save_sheet(
            [
                ["Date", "Value"],
                ["2024-01-31", 5.0],
                ["2024-02-29", 6.0],
            ]
        )
        records = melt_sheet(content, "Sheet1")
        by_date = {r["date"]: r["value"] for r in records}
        assert by_date[date(2024, 1, 31)] == 5.0
        assert by_date[date(2024, 2, 29)] == 6.0

    def test_date_bounds_filter_rows(self):
        """Inclusive start and end bounds drop out-of-range observations."""
        content = _save_sheet(
            [
                ["Date", "Value"],
                [datetime(2023, 1, 31), 1.0],
                [datetime(2024, 1, 31), 2.0],
                [datetime(2025, 1, 31), 3.0],
            ]
        )
        records = melt_sheet(
            content,
            "Sheet1",
            start_date=date(2024, 1, 1),
            end_date=date(2024, 12, 31),
        )
        assert [r["date"] for r in records] == [date(2024, 1, 31)]
