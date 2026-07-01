"""Tests for the additional New York Fed regional dataset models."""

import io
from datetime import date
from unittest.mock import MagicMock

import pytest
from openbb_core.provider.utils.errors import EmptyDataError

from openbb_federal_reserve.models.regional.new_york_empire import (
    FederalReserveNewYorkEmpireStateData,
    FederalReserveNewYorkEmpireStateFetcher,
)
from openbb_federal_reserve.models.regional.new_york_hhdc import (
    FederalReserveNewYorkHouseholdDebtData,
    FederalReserveNewYorkHouseholdDebtFetcher,
)
from openbb_federal_reserve.models.regional.new_york_mct import (
    FederalReserveNewYorkCoreTrendInflationData,
    FederalReserveNewYorkCoreTrendInflationFetcher,
)
from openbb_federal_reserve.models.regional.new_york_sce_credit import (
    FederalReserveNewYorkConsumerCreditAccessData,
    FederalReserveNewYorkConsumerCreditAccessFetcher,
)
from openbb_federal_reserve.models.regional.new_york_sce_housing import (
    FederalReserveNewYorkConsumerHousingData,
    FederalReserveNewYorkConsumerHousingFetcher,
)
from openbb_federal_reserve.models.regional.new_york_sce_labor import (
    FederalReserveNewYorkConsumerLaborMarketData,
    FederalReserveNewYorkConsumerLaborMarketFetcher,
)

_EMPIRE_OVERVIEW = (
    '<a href="/medialibrary/media/survey/empire/data/'
    'esms_seasonallyadjusted_diffusion.csv?sc_lang=en&amp;hash=ABC">a</a>'
    '<a href="/medialibrary/media/survey/empire/data/'
    'esms_seasonallyadjusted_allseries.csv?sc_lang=en&amp;hash=DEF">b</a>'
)
_EMPIRE_DIFFUSION_CSV = (
    "surveyDate,GACDISA,NOCDISA,GAFDISA\n"
    "2026-05-31,19.6,1.0,25.0\n"
    "2026-06-30,5.7,ND,30.1\n"
)
_EMPIRE_ALLSERIES_CSV = (
    "surveyDate,GACDSA,GACISA,GACNSA,GACDISA\n2026-06-30,20.0,25.7,54.3,5.7\n"
)


def _response(content: bytes) -> MagicMock:
    """Build a make_request response with the given binary body."""
    response = MagicMock()
    response.content = content
    response.raise_for_status = MagicMock()
    return response


def _text_response(text: str) -> MagicMock:
    """Build a make_request response with the given text body."""
    response = MagicMock()
    response.text = text
    response.content = text.encode("utf-8")
    response.raise_for_status = MagicMock()
    return response


def _empire_request(url, *args, **kwargs):
    """Path-aware fake: overview HTML, then the diffusion or all-series CSV."""
    if url.endswith("empiresurvey_overview"):
        return _text_response(_EMPIRE_OVERVIEW)
    if "allseries" in url:
        return _text_response(_EMPIRE_ALLSERIES_CSV)
    return _text_response(_EMPIRE_DIFFUSION_CSV)


def _labor_workbook() -> bytes:
    """Build an SCE Labor Market workbook with the five-row header offset."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for _ in range(5):
        sheet.append([None])
    sheet.append(["Date", "Mean", "Dispersion"])
    sheet.append([None, None, None])
    sheet.append(["Mar 2025", 84.0, "[40,72]"])
    sheet.append(["Jul 2025", 85.5, "[41,73]"])
    sheet.append(["Mar 2026", 86.0, "[42,74]"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _housing_workbook() -> bytes:
    """Build an SCE Housing workbook with the four-row header offset."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Home Price Expectations"
    for _ in range(4):
        sheet.append([None])
    sheet.append(
        [
            "date",
            "Average One-Year Ahead Home Price Change Expectations",
            "Median One-Year Ahead Home Price Change Expectations",
            "25th Percentile one-year ahead Home Price Change Expectations",
            "75th Percentile one-year ahead Home Price Change Expectations",
            "Average Five-Year Ahead Home Price Change Expectations",
            "Median Five-Year Ahead Home Price Change Expectations",
        ]
    )
    sheet.append([202502, 5.6, 4.0, 0.0, 10.0, 3.3, 3.0])
    sheet.append([202602, 5.2, 4.0, 0.0, 10.0, 3.2, 3.0])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _credit_workbook() -> bytes:
    """Build an SCE Credit Access workbook with overall and demographics sheets."""
    from openpyxl import Workbook

    header = [
        "date",
        "group",
        "category",
        "Applied_Accepted",
        "Applied_Rejected",
        "Discouraged",
    ]
    workbook = Workbook()
    overall = workbook.active
    overall.title = "overall"
    overall.append([None])
    overall.append(header)
    overall.append([202510, "all", "Overall", 34.0, 6.0, 8.0])
    overall.append([202602, "all", "Overall", 33.971748, 6.431378, 8.265976])
    demographics = workbook.create_sheet("demographics")
    demographics.append([None])
    demographics.append(header)
    demographics.append([202602, "credit_score", "less_680", 20.0, 18.0, 15.0])
    demographics.append([202602, "credit_score", "over_760", 45.0, 2.0, 3.0])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _hhdc_workbook() -> bytes:
    """Build an HHDC workbook with rows-period, cols-period, and snapshot sheets."""
    from openpyxl import Workbook

    workbook = Workbook()
    toc = workbook.active
    toc.title = "TABLE OF CONTENTS"
    toc.append([None, "Total Debt Balance and its Composition", "Page 3"])
    toc.append([None, "Debt Share by Product Type and Age (2026Q1)", "Page 21"])
    toc.append([None, "Total Debt Balance per Capita* by State", "Page 32"])

    page3 = workbook.create_sheet("Page 3 Data")
    page3.append(["Total Debt Balance and its Composition"])
    page3.append(["Trillions of $"])
    page3.append(["Return to Table of Contents"])
    page3.append(
        [
            None,
            "Mortgage",
            "HE Revolving",
            "Auto Loan",
            "Credit Card",
            "Student Loan",
            "Other",
            "Total",
        ]
    )
    page3.append(["25:Q4", 13.17, 0.43, 1.66, 1.27, 1.66, 0.56, 18.77])
    page3.append(["26:Q1", 13.19, 0.44, 1.68, 1.25, 1.65, 0.56, 18.79])

    page21 = workbook.create_sheet("Page 21 Data")
    page21.append(["Debt Share by Product Type and Age (2026Q1)"])
    page21.append(["Trillions of Dollars"])
    page21.append([None, "18-29", "30-39", "40-49"])
    page21.append(["Auto Loans", 0.19, 0.37, 0.40])
    page21.append(["Credit Card", 0.07, 0.21, 0.29])

    page32 = workbook.create_sheet("Page 32 Data")
    page32.append(["Total Debt Balance per Capita* by State"])
    page32.append(["Thousands of $"])
    page32.append([None, "25:Q4", "26:Q1"])
    page32.append(["AZ", 73.78, 73.06])
    page32.append(["CA", 88.37, 87.71])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


_MCT_CSV = (
    "Section,,,,,,,,,,,,,,,,\n"
    "Route,,,,,,,,,,,,,,,,\n"
    "Radio,,,,,,,,,,,,,,,,\n"
    "column name,Date,MCT,MCT,MCT,MCT,Headline PCE inflation (YoY),"
    "Core PCE inflation (YoY),Goods,Services ex. housing,Housing,"
    "Goods: Common,Goods: Sector-specific,Services ex. housing: Common,"
    "Services ex. housing: Sector-specific,Housing: Common,"
    "Housing: Sector-specific\n"
    ",3/1/2026,3.11,3.49,3.88,1.66,3.6,3.2,0.8,0.9,0.4,"
    "0.46,0.23,0.52,0.2,0.31,-0.13\n"
    ",4/1/2026,3.52,4.02,4.50,2.19,3.77,3.29,0.87,0.89,0.44,"
    "0.61,0.23,0.68,0.17,0.43,-0.03\n"
)


class TestEmpireStateManufacturing:
    """Tests for the Empire State Manufacturing fetcher."""

    def test_scrapes_hash_melts_and_filters(self, monkeypatch):
        """The diffusion CSV scrapes, pivots to wide, ND coerces, dates filter."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request", _empire_request
        )
        query = FederalReserveNewYorkEmpireStateFetcher.transform_query(
            {"start_date": "2026-06-01", "end_date": "2026-06-30"}
        )
        rows = FederalReserveNewYorkEmpireStateFetcher.extract_data(query, None)
        result = FederalReserveNewYorkEmpireStateFetcher.transform_data(query, rows)
        assert all(isinstance(r, FederalReserveNewYorkEmpireStateData) for r in result)
        assert len(result) == 1
        assert result[0].date == date(2026, 6, 30)
        row = result[0].model_dump()
        assert row["general_business_conditions_current_diffusion_index"] == 5.7
        assert row["new_orders_current_diffusion_index"] is None
        assert row["general_business_conditions_future_diffusion_index"] == 30.1

    def test_all_series_dataset_exposes_response_shares(self, monkeypatch):
        """The all-series dataset exposes the up / down / same response shares."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request", _empire_request
        )
        query = FederalReserveNewYorkEmpireStateFetcher.transform_query(
            {"dataset": "seasonally_adjusted_all_series"}
        )
        rows = FederalReserveNewYorkEmpireStateFetcher.extract_data(query, None)
        result = FederalReserveNewYorkEmpireStateFetcher.transform_data(query, rows)
        assert len(result) == 1
        row = result[0].model_dump()
        assert row["general_business_conditions_current_share_increase"] == 25.7
        assert row["general_business_conditions_current_share_decrease"] == 20.0
        assert row["general_business_conditions_current_share_same"] == 54.3
        assert row["general_business_conditions_current_diffusion_index"] == 5.7

    def test_empty_overview_raises(self, monkeypatch):
        """An overview page without the CSV link raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _text_response("<html>no link</html>"),
        )
        query = FederalReserveNewYorkEmpireStateFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkEmpireStateFetcher.extract_data(query, None)


class TestConsumerLaborMarket:
    """Tests for the SCE Labor Market fetcher."""

    def test_parses_mean_and_filters(self, monkeypatch):
        """The earnings sheet parses the mean reservation wage and dates filter."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_labor_workbook()),
        )
        query = FederalReserveNewYorkConsumerLaborMarketFetcher.transform_query(
            {"start_date": "2026-01-01", "end_date": "2026-03-31"}
        )
        rows = FederalReserveNewYorkConsumerLaborMarketFetcher.extract_data(query, None)
        result = FederalReserveNewYorkConsumerLaborMarketFetcher.transform_data(
            query, rows
        )
        assert all(
            isinstance(r, FederalReserveNewYorkConsumerLaborMarketData) for r in result
        )
        assert len(result) == 1
        assert result[0].date == date(2026, 3, 1)
        assert result[0].average_reservation_wage == 86.0 * 1000

    def test_empty_raises(self, monkeypatch):
        """An empty response raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(b""),
        )
        query = FederalReserveNewYorkConsumerLaborMarketFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkConsumerLaborMarketFetcher.extract_data(query, None)


class TestConsumerHousing:
    """Tests for the SCE Housing fetcher."""

    def test_parses_yyyymm_and_filters(self, monkeypatch):
        """The home-price sheet pivots to wide rows; YYYYMM dates and filters apply."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_housing_workbook()),
        )
        query = FederalReserveNewYorkConsumerHousingFetcher.transform_query(
            {"start_date": "2025-01-01", "end_date": "2025-12-31"}
        )
        rows = FederalReserveNewYorkConsumerHousingFetcher.extract_data(query, None)
        result = FederalReserveNewYorkConsumerHousingFetcher.transform_data(query, rows)
        assert all(
            isinstance(r, FederalReserveNewYorkConsumerHousingData) for r in result
        )
        assert len(result) == 1
        assert result[0].date == date(2025, 2, 1)
        row = result[0].model_dump()
        assert row["Average One-Year Ahead Home Price Change Expectations"] == 5.6
        assert row["Median Five-Year Ahead Home Price Change Expectations"] == 3.0

    def test_empty_raises(self, monkeypatch):
        """An empty response raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(b""),
        )
        query = FederalReserveNewYorkConsumerHousingFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkConsumerHousingFetcher.extract_data(query, None)


class TestConsumerCreditAccess:
    """Tests for the SCE Credit Access fetcher."""

    def test_parses_overall_and_filters(self, monkeypatch):
        """The overall sheet pivots every metric to wide rows and filters."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_credit_workbook()),
        )
        query = FederalReserveNewYorkConsumerCreditAccessFetcher.transform_query(
            {"start_date": "2026-01-01", "end_date": "2026-02-28"}
        )
        rows = FederalReserveNewYorkConsumerCreditAccessFetcher.extract_data(
            query, None
        )
        result = FederalReserveNewYorkConsumerCreditAccessFetcher.transform_data(
            query, rows
        )
        assert all(
            isinstance(r, FederalReserveNewYorkConsumerCreditAccessData) for r in result
        )
        assert len(result) == 1
        assert result[0].date == date(2026, 2, 1)
        assert result[0].group == "all"
        row = result[0].model_dump()
        assert row["Applied and Accepted"] == 33.971748
        assert row["Applied and Rejected"] == 6.431378
        assert row["Discouraged from Applying"] == 8.265976

    def test_demographics_breakdown(self, monkeypatch):
        """The demographics sheet exposes the group and category dimensions."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_credit_workbook()),
        )
        query = FederalReserveNewYorkConsumerCreditAccessFetcher.transform_query(
            {"breakdown": "demographics"}
        )
        rows = FederalReserveNewYorkConsumerCreditAccessFetcher.extract_data(
            query, None
        )
        result = FederalReserveNewYorkConsumerCreditAccessFetcher.transform_data(
            query, rows
        )
        assert all(r.group == "credit_score" for r in result)
        by_category = {r.category: r.model_dump() for r in result}
        assert by_category["less_680"]["Applied and Accepted"] == 20.0
        assert by_category["over_760"]["Applied and Rejected"] == 2.0

    def test_empty_raises(self, monkeypatch):
        """An empty response raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(b""),
        )
        query = FederalReserveNewYorkConsumerCreditAccessFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkConsumerCreditAccessFetcher.extract_data(query, None)


class TestHouseholdDebt:
    """Tests for the Household Debt and Credit fetcher."""

    def test_discovers_quarter_melts_and_filters(self, monkeypatch):
        """The latest quarter discovers, the default table pivots, and dates filter."""
        from openbb_federal_reserve.utils import ny_hhdc

        monkeypatch.setattr(ny_hhdc, "latest_household_debt_quarter", lambda: "2026Q1")
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_hhdc_workbook()),
        )
        query = FederalReserveNewYorkHouseholdDebtFetcher.transform_query(
            {"start_date": "2026-01-01", "end_date": "2026-03-31"}
        )
        rows = FederalReserveNewYorkHouseholdDebtFetcher.extract_data(query, None)
        result = FederalReserveNewYorkHouseholdDebtFetcher.transform_data(query, rows)
        assert all(
            isinstance(r, FederalReserveNewYorkHouseholdDebtData) for r in result
        )
        assert len(result) == 1
        assert result[0].date == date(2026, 3, 31)
        row = result[0].model_dump()
        assert row["Mortgage"] == 13.19
        assert row["Total"] == 18.79

    def test_by_state_table_melts_columns_as_periods(self, monkeypatch):
        """The by-state table reads periods from the header row and states as rows."""
        from openbb_federal_reserve.utils import ny_hhdc

        monkeypatch.setattr(ny_hhdc, "latest_household_debt_quarter", lambda: "2026Q1")
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_hhdc_workbook()),
        )
        query = FederalReserveNewYorkHouseholdDebtFetcher.transform_query(
            {"table": "composition_of_debt_balance_per_capita_by_state"}
        )
        rows = FederalReserveNewYorkHouseholdDebtFetcher.extract_data(query, None)
        result = FederalReserveNewYorkHouseholdDebtFetcher.transform_data(query, rows)
        by_date = {str(r.date): r.model_dump() for r in result}
        assert by_date["2026-03-31"]["AZ"] == 73.06
        assert by_date["2025-12-31"]["CA"] == 88.37

    def test_snapshot_table_uses_title_quarter(self, monkeypatch):
        """The cross-section snapshot table dates every row from the title quarter."""
        from openbb_federal_reserve.utils import ny_hhdc

        monkeypatch.setattr(ny_hhdc, "latest_household_debt_quarter", lambda: "2026Q1")
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_hhdc_workbook()),
        )
        query = FederalReserveNewYorkHouseholdDebtFetcher.transform_query(
            {"table": "debt_share_by_product_type_and_age"}
        )
        rows = FederalReserveNewYorkHouseholdDebtFetcher.extract_data(query, None)
        result = FederalReserveNewYorkHouseholdDebtFetcher.transform_data(query, rows)
        assert len(result) == 1
        assert result[0].date == date(2026, 3, 31)
        row = result[0].model_dump()
        assert row["Auto Loans - 18-29"] == 0.19
        assert row["Credit Card - 40-49"] == 0.29

    def test_empty_raises(self, monkeypatch):
        """An empty response raises ``EmptyDataError``."""
        from openbb_federal_reserve.utils import ny_hhdc

        monkeypatch.setattr(ny_hhdc, "latest_household_debt_quarter", lambda: "2026Q1")
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(b""),
        )
        query = FederalReserveNewYorkHouseholdDebtFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkHouseholdDebtFetcher.extract_data(query, None)


class TestCoreTrendInflation:
    """Tests for the Multivariate Core Trend Inflation fetcher."""

    def test_parses_by_position_and_filters(self, monkeypatch):
        """The CSV parses by column position past its metadata rows and dates filter."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _text_response(_MCT_CSV),
        )
        query = FederalReserveNewYorkCoreTrendInflationFetcher.transform_query(
            {"start_date": "2026-04-01", "end_date": "2026-04-30"}
        )
        rows = FederalReserveNewYorkCoreTrendInflationFetcher.extract_data(query, None)
        result = FederalReserveNewYorkCoreTrendInflationFetcher.transform_data(
            query, rows
        )
        assert all(
            isinstance(r, FederalReserveNewYorkCoreTrendInflationData) for r in result
        )
        assert len(result) == 1
        assert result[0].date == date(2026, 4, 1)
        assert result[0].mct == 3.52
        assert result[0].mct_lower == 2.19
        assert result[0].headline_pce == 3.77
        assert result[0].housing == 0.44
        assert result[0].goods_common == 0.61
        assert result[0].goods_sector_specific == 0.23
        assert result[0].housing_sector_specific == -0.03

    def test_empty_raises(self, monkeypatch):
        """An empty response raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _text_response(""),
        )
        query = FederalReserveNewYorkCoreTrendInflationFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkCoreTrendInflationFetcher.extract_data(query, None)


def _sce_workbook(title: str, rows: list[list]) -> bytes:
    """Build a one-sheet SCE workbook from explicit rows."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class TestParseSce:
    """Direct tests for the SCE workbook melter."""

    def test_too_few_rows_returns_empty(self):
        """A sheet without a data row below the header melts to nothing."""
        from openbb_federal_reserve.utils.sce import parse_sce

        content = _sce_workbook("Tiny", [["title"], [None], [None], [None, "Mean"]])
        assert parse_sce(content, "Tiny") == []

    def test_group_label_joins_onto_column_label(self):
        """A forward-filled group label distinct from the column label joins onto it."""
        from openbb_federal_reserve.utils.sce import parse_sce

        content = _sce_workbook(
            "Grouped",
            [
                ["title"],
                [None],
                [None, "Year ahead", None],
                [None, "Mean", "Median"],
                [202602, 3.0, 2.9],
            ],
        )
        result = parse_sce(content, "Grouped")
        by_series = {row["series"]: row["value"] for row in result}
        assert by_series["Year ahead - Mean"] == 3.0
        assert by_series["Year ahead - Median"] == 2.9

    def test_demographic_blocks_disambiguate_and_skip_bad_dates(self):
        """Repeated labels across blank-separated blocks gain ``Estimate`` prefixes."""
        from openbb_federal_reserve.utils.sce import parse_sce

        content = _sce_workbook(
            "Demo",
            [
                ["title"],
                [None, "Age < 40", "Age < 40", None, "Age >= 40", "Age >= 40"],
                [None, "Mean", "Median", None, "Mean", "Median"],
                [None, "Mean", "Median", None, "Mean", "Median"],
                [202601, 3.0, 2.9, None, 4.0, 3.8],
                ["not-a-date", 1.0, 1.0, None, 1.0, 1.0],
            ],
        )
        result = parse_sce(content, "Demo")
        by_series = {row["series"]: row["value"] for row in result}
        assert by_series["Estimate 1 - Mean"] == 3.0
        assert by_series["Estimate 2 - Median"] == 3.8
        assert all(row["date"] == date(2026, 1, 1) for row in result)

    def test_start_date_filters_earlier_rows(self):
        """A ``start_date`` drops observations before the window."""
        from openbb_federal_reserve.utils.sce import parse_sce

        content = _sce_workbook(
            "Range",
            [
                ["title"],
                [None, "Mean"],
                [None, "Mean"],
                [None, "Mean"],
                [202412, 1.0],
                [202602, 2.0],
            ],
        )
        result = parse_sce(content, "Range", start_date=date(2026, 1, 1))
        assert len(result) == 1
        assert result[0]["date"] == date(2026, 2, 1)
        assert result[0]["value"] == 2.0


def _hhdc_page_workbook(page: int, rows: list[list]) -> bytes:
    """Build a single ``Page N Data`` HHDC workbook from explicit rows."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"Page {page} Data"
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class TestParseHhdcSheet:
    """Direct tests for the HHDC sheet melter and its date helpers."""

    def test_coerce_date_handles_datetime_and_non_quarter(self):
        """``_coerce_date`` unwraps a datetime and rejects non-date inputs."""
        import datetime as _dt

        from openbb_federal_reserve.utils.ny_hhdc import _coerce_date

        assert _coerce_date(_dt.datetime(2026, 3, 31, 12)) == date(2026, 3, 31)
        assert _coerce_date(_dt.date(2026, 3, 31)) == date(2026, 3, 31)
        assert _coerce_date(123) is None

    def test_title_quarter_without_quarter_returns_none(self):
        """A title with no ``(YYYY Qn)`` token yields no date."""
        from openbb_federal_reserve.utils.ny_hhdc import _title_quarter

        assert _title_quarter("No quarter here") is None

    def test_missing_sheet_returns_empty(self):
        """A workbook missing the requested page melts to nothing."""
        from openbb_federal_reserve.utils.ny_hhdc import parse_hhdc_sheet

        content = _hhdc_page_workbook(3, [["x"]])
        assert parse_hhdc_sheet(content, 99) == []

    def test_period_in_rows_skips_bad_cells_and_dates(self):
        """Non-numeric cells, blank cells, and unparseable date rows are dropped."""
        from openbb_federal_reserve.utils.ny_hhdc import parse_hhdc_sheet

        content = _hhdc_page_workbook(
            3,
            [
                ["Total Debt Balance and its Composition"],
                ["Trillions"],
                ["Return to Table of Contents"],
                [None, "Mortgage", "Auto"],
                ["25:Q4", 13.1, "notanumber"],
                ["notaquarter", 1.0, 2.0],
                ["26:Q1", None, 1.7],
            ],
        )
        result = parse_hhdc_sheet(content, 3)
        by_key = {(str(r["date"]), r["series"]): r["value"] for r in result}
        assert by_key[("2025-12-31", "Mortgage")] == 13.1
        assert by_key[("2026-03-31", "Auto")] == 1.7
        # The non-numeric and blank cells produce no record, and the
        # unparseable date row is skipped entirely.
        assert ("2025-12-31", "Auto") not in by_key
        assert ("2026-03-31", "Mortgage") not in by_key

    def test_period_in_columns_skips_non_string_categories(self):
        """In the by-state layout, non-string row labels are skipped."""
        from openbb_federal_reserve.utils.ny_hhdc import parse_hhdc_sheet

        content = _hhdc_page_workbook(
            32,
            [
                ["Per Capita by State"],
                ["Thousands"],
                [None, "25:Q4", "26:Q1"],
                ["AZ", 73.8, 73.0],
                [123, 1.0, 2.0],
                ["   ", 1.0, 2.0],
            ],
        )
        result = parse_hhdc_sheet(content, 32)
        assert {r["series"] for r in result} == {"AZ"}

    def test_snapshot_falls_back_to_default_dimension_row(self):
        """A snapshot with under two labels per row uses the default dimension row."""
        from openbb_federal_reserve.utils.ny_hhdc import parse_hhdc_sheet

        content = _hhdc_page_workbook(
            21,
            [
                ["Debt Share (2026Q1)"],
                ["Trillions"],
                [None, "18-29"],
                ["Auto", 0.19],
                [99, 1.0],
            ],
        )
        result = parse_hhdc_sheet(content, 21)
        by_series = {r["series"]: r["value"] for r in result}
        assert by_series["Auto - 18-29"] == 0.19
        assert all(r["date"] == date(2026, 3, 31) for r in result)


class TestDecodeEmpireColumn:
    """Direct tests for the Empire State column-code decoder."""

    def test_undecodable_columns_are_dropped(self):
        """Columns with a bad adjustment suffix or unknown indicator are ignored."""
        from openbb_federal_reserve.utils.ny_empire_survey import parse

        text = (
            "surveyDate,GACDISA,ZZCDISA,JUNK\n"
            "2026-05-31,1.0,2.0,3.0\n"
            "2026-06-30,5.7,2.0,3.0\n"
        )
        result = parse(text, start_date=date(2026, 6, 1), end_date=date(2026, 6, 30))
        series = {row["series"] for row in result}
        assert series == {"general_business_conditions_current_diffusion_index"}
        assert all(row["date"] == date(2026, 6, 30) for row in result)
