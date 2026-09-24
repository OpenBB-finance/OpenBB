"""Tests for the New York Fed regional models."""

import io
from datetime import date
from unittest.mock import MagicMock

import pytest
from openbb_core.provider.utils.errors import EmptyDataError

from openbb_federal_reserve.models.regional.new_york_bls import (
    FederalReserveNewYorkBusinessLeadersData,
    FederalReserveNewYorkBusinessLeadersFetcher,
)
from openbb_federal_reserve.models.regional.new_york_cmdi import (
    FederalReserveNewYorkCorporateDistressData,
    FederalReserveNewYorkCorporateDistressFetcher,
)
from openbb_federal_reserve.models.regional.new_york_gscpi import (
    FederalReserveNewYorkSupplyChainData,
    FederalReserveNewYorkSupplyChainFetcher,
)
from openbb_federal_reserve.models.regional.new_york_nowcast import (
    FederalReserveNewYorkNowcastData,
    FederalReserveNewYorkNowcastFetcher,
)
from openbb_federal_reserve.models.regional.new_york_sce import (
    FederalReserveNewYorkConsumerExpectationsData,
    FederalReserveNewYorkConsumerExpectationsFetcher,
)


def _cmdi_workbook() -> bytes:
    """Build a CMDI workbook with the five-row header offset."""
    from datetime import datetime

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Index Data"
    for _ in range(5):
        sheet.append([None])
    sheet.append(["eow_friday", "Market CMDI", "IG CMDI", "HY CMDI"])
    sheet.append([datetime(2026, 6, 12), 0.12, 0.22, 0.05])
    sheet.append([datetime(2026, 6, 19), 0.13, 0.23, 0.06])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sce_workbook() -> bytes:
    """Build an SCE inflation-expectations workbook with a three-row offset."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inflation expectations"
    for _ in range(3):
        sheet.append([None])
    sheet.append(
        [
            "date",
            "Median one-year ahead expected inflation rate",
            "Median three-year ahead expected inflation rate",
            "25th Percentile one-year ahead expected inflation rate",
            "75th Percentile one-year ahead expected inflation rate",
            "25th Percentile three-year ahead expected inflation rate",
            "75th Percentile three-year ahead expected inflation rate",
        ]
    )
    sheet.append([202604, 3.40, 3.10, 2.00, 6.00, 1.00, 5.90])
    sheet.append([202605, 3.46, 3.13, 1.99, 6.00, 1.00, 5.94])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sce_credit_workbook() -> bytes:
    """Build an SCE credit-availability sheet with a two-level header."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Credit availability"
    sheet.append([None])
    sheet.append([None])
    sheet.append([None, "Year ago", None, "Year ahead"])
    sheet.append([None, "Much harder", "Much easier", "Much harder"])
    sheet.append([202605, 10.0, 20.0, 5.0])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _gscpi_workbook() -> bytes:
    """Build a GSCPI workbook with a branding-noise row to drop."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "GSCPI Monthly Data"
    sheet.append(["Date", "GSCPI", None, None])
    sheet.append([None, None, None, "NEW YORK FED ECONOMIC RESEARCH"])
    sheet.append(["30-Apr-2026", 1.8231, None, None])
    sheet.append(["31-May-2026", 1.7689, None, None])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _nowcast_workbook() -> bytes:
    """Build a Nowcast workbook with the five-row branding offset."""
    from datetime import datetime

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Forecasts By Horizon"
    for _ in range(5):
        sheet.append([None])
    sheet.append(
        [
            "Forecast date",
            "Reference quarter",
            "Backcast (previous quarter)",
            "Nowcast (current quarter)",
            "Forecast (next quarter)",
        ]
    )
    sheet.append([datetime(2026, 6, 19), "2026:Q2", None, 2.68, 2.45])
    sheet.append([datetime(2026, 6, 26), "2026:Q2", None, 2.70, 2.42])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _response(content: bytes) -> MagicMock:
    """Build a make_request response with the given binary body."""
    response = MagicMock()
    response.content = content
    response.raise_for_status = MagicMock()
    return response


def _text_response(text: str) -> MagicMock:
    """Build a make_request response with the given text body."""
    response = MagicMock()
    response.text = text
    response.raise_for_status = MagicMock()
    return response


_BLS = "Dates,Current,Expected\n05/1/2026,-8.0,10.0\n06/1/2026,-10.1,8.3\n"


class TestBusinessLeaders:
    """Tests for the Business Leaders Survey fetcher."""

    def test_parses_and_filters(self, monkeypatch):
        """The survey CSV parses and dates filter."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _text_response(_BLS),
        )
        query = FederalReserveNewYorkBusinessLeadersFetcher.transform_query(
            {"start_date": "2026-06-01", "end_date": "2026-06-30"}
        )
        rows = FederalReserveNewYorkBusinessLeadersFetcher.extract_data(query, None)
        result = FederalReserveNewYorkBusinessLeadersFetcher.transform_data(query, rows)
        assert all(
            isinstance(r, FederalReserveNewYorkBusinessLeadersData) for r in result
        )
        assert len(result) == 1
        assert result[0].date == date(2026, 6, 1)
        assert result[0].current == -10.1
        assert result[0].expected == 8.3

    def test_empty_raises(self, monkeypatch):
        """An empty response raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _text_response(""),
        )
        query = FederalReserveNewYorkBusinessLeadersFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkBusinessLeadersFetcher.extract_data(query, None)


class TestSupplyChainPressure:
    """Tests for the GSCPI fetcher."""

    def test_parses_drops_noise_and_filters(self, monkeypatch):
        """The monthly sheet parses, branding rows drop, and dates filter."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_gscpi_workbook()),
        )
        query = FederalReserveNewYorkSupplyChainFetcher.transform_query(
            {"start_date": "2026-05-01", "end_date": "2026-05-31"}
        )
        rows = FederalReserveNewYorkSupplyChainFetcher.extract_data(query, None)
        result = FederalReserveNewYorkSupplyChainFetcher.transform_data(query, rows)
        assert all(isinstance(r, FederalReserveNewYorkSupplyChainData) for r in result)
        assert len(result) == 1
        assert result[0].date == date(2026, 5, 31)
        assert result[0].gscpi == 1.7689

    def test_empty_raises(self, monkeypatch):
        """An empty response raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(b""),
        )
        query = FederalReserveNewYorkSupplyChainFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkSupplyChainFetcher.extract_data(query, None)


class TestNowcast:
    """Tests for the New York Fed Staff Nowcast fetcher."""

    def test_parses_by_horizon(self, monkeypatch):
        """The by-horizon sheet parses past its offset with typed values."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_nowcast_workbook()),
        )
        query = FederalReserveNewYorkNowcastFetcher.transform_query({})
        rows = FederalReserveNewYorkNowcastFetcher.extract_data(query, None)
        result = FederalReserveNewYorkNowcastFetcher.transform_data(query, rows)
        assert all(isinstance(r, FederalReserveNewYorkNowcastData) for r in result)
        assert [r.date for r in result] == [date(2026, 6, 19), date(2026, 6, 26)]
        assert result[0].reference_quarter == "2026:Q2"
        assert result[0].backcast is None
        assert result[0].nowcast == 2.68
        assert result[1].forecast == 2.42

    def test_date_filters(self, monkeypatch):
        """The start and end date filters narrow the vintages."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_nowcast_workbook()),
        )
        query = FederalReserveNewYorkNowcastFetcher.transform_query(
            {"start_date": "2026-06-20", "end_date": "2026-06-30"}
        )
        rows = FederalReserveNewYorkNowcastFetcher.extract_data(query, None)
        result = FederalReserveNewYorkNowcastFetcher.transform_data(query, rows)
        assert [r.date for r in result] == [date(2026, 6, 26)]

    def test_empty_content_raises(self, monkeypatch):
        """An empty response raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(b""),
        )
        query = FederalReserveNewYorkNowcastFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkNowcastFetcher.extract_data(query, None)

    def test_no_rows_after_filter_raises(self, monkeypatch):
        """A filter that removes every vintage raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_nowcast_workbook()),
        )
        query = FederalReserveNewYorkNowcastFetcher.transform_query(
            {"start_date": "2030-01-01"}
        )
        rows = FederalReserveNewYorkNowcastFetcher.extract_data(query, None)
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkNowcastFetcher.transform_data(query, rows)


class TestCorporateBondDistress:
    """Tests for the CMDI fetcher."""

    def test_parses_index_sheet_and_filters(self, monkeypatch):
        """The CMDI index sheet parses past its header offset and dates filter."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_cmdi_workbook()),
        )
        query = FederalReserveNewYorkCorporateDistressFetcher.transform_query(
            {"start_date": "2026-06-15", "end_date": "2026-06-30"}
        )
        rows = FederalReserveNewYorkCorporateDistressFetcher.extract_data(query, None)
        result = FederalReserveNewYorkCorporateDistressFetcher.transform_data(
            query, rows
        )
        assert all(
            isinstance(r, FederalReserveNewYorkCorporateDistressData) for r in result
        )
        assert len(result) == 1
        assert result[0].date == date(2026, 6, 19)
        assert result[0].market == 0.13
        assert result[0].high_yield == 0.06

    def test_empty_raises(self, monkeypatch):
        """An empty response raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(b""),
        )
        query = FederalReserveNewYorkCorporateDistressFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkCorporateDistressFetcher.extract_data(query, None)


class TestConsumerExpectations:
    """Tests for the SCE fetcher."""

    def test_parses_yyyymm_dates_and_filters(self, monkeypatch):
        """The topic sheet pivots to wide rows; YYYYMM dates and filters apply."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_sce_workbook()),
        )
        query = FederalReserveNewYorkConsumerExpectationsFetcher.transform_query(
            {"start_date": "2026-01-01", "end_date": "2026-04-30"}
        )
        rows = FederalReserveNewYorkConsumerExpectationsFetcher.extract_data(
            query, None
        )
        result = FederalReserveNewYorkConsumerExpectationsFetcher.transform_data(
            query, rows
        )
        assert all(
            isinstance(r, FederalReserveNewYorkConsumerExpectationsData) for r in result
        )
        assert len(result) == 1
        assert result[0].date == date(2026, 4, 1)
        row = result[0].model_dump()
        assert row["Median one-year ahead expected inflation rate"] == 3.40
        assert row["Median three-year ahead expected inflation rate"] == 3.10

    def test_group_prefixed_series(self, monkeypatch):
        """A two-level header joins the forward-filled group onto each label."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(_sce_credit_workbook()),
        )
        query = FederalReserveNewYorkConsumerExpectationsFetcher.transform_query(
            {"topic": "credit_availability"}
        )
        rows = FederalReserveNewYorkConsumerExpectationsFetcher.extract_data(
            query, None
        )
        result = FederalReserveNewYorkConsumerExpectationsFetcher.transform_data(
            query, rows
        )
        assert len(result) == 1
        row = result[0].model_dump()
        assert row["Year ago - Much harder"] == 10.0
        assert row["Year ago - Much easier"] == 20.0
        assert row["Year ahead - Much harder"] == 5.0

    def test_empty_raises(self, monkeypatch):
        """An empty response raises ``EmptyDataError``."""
        monkeypatch.setattr(
            "openbb_core.provider.utils.helpers.make_request",
            lambda *a, **k: _response(b""),
        )
        query = FederalReserveNewYorkConsumerExpectationsFetcher.transform_query({})
        with pytest.raises(EmptyDataError):
            FederalReserveNewYorkConsumerExpectationsFetcher.extract_data(query, None)


class TestPrimaryDealerFailsWidget:
    """Tests for the ``primary_dealer_fails`` widget column ordering."""

    @staticmethod
    def _table() -> dict:
        """Return the ``primary_dealer_fails`` widget table config."""
        from openbb_federal_reserve.regional import new_york

        route = next(
            r
            for r in new_york.router.api_router.routes
            if getattr(r, "name", "") == "primary_dealer_fails"
        )
        return route.openapi_extra["widget_config"]["data"]["table"]  # ty: ignore[unresolved-attribute]

    def test_pins_ftd_then_ftr_groups_in_order(self):
        """FTD columns are pinned together, then FTR, in protocol asset order."""
        columns = self._table()["columnsDefs"]
        assert columns[0]["field"] == "date"
        ordered = [c["field"] for c in columns[1:]]
        expected = [
            "FTD Treasury Securities (Ex-TIPS)",
            "FTD TIPS",
            "FTD Agency and GSE Securities (Ex-MBS)",
            "FTD Agency and GSE MBS",
            "FTD Other MBS",
            "FTD Corporate Securities",
            "FTD Total",
            "FTR Treasury Securities (Ex-TIPS)",
            "FTR TIPS",
            "FTR Agency and GSE Securities (Ex-MBS)",
            "FTR Agency and GSE MBS",
            "FTR Other MBS",
            "FTR Corporate Securities",
            "FTR Total",
        ]
        assert ordered == expected
        ftd = [f for f in ordered if f.startswith("FTD")]
        ftr = [f for f in ordered if f.startswith("FTR")]
        assert ordered == ftd + ftr

    def test_pinned_fields_match_model_titles_and_show_all_kept(self):
        """Pinned non-total fields are real model titles; ``showAll`` stays on."""
        from openbb_federal_reserve.utils.primary_dealer_statistics import (
            FAILS_SERIES_TO_TITLE,
        )

        table = self._table()
        assert table["showAll"] is True
        model_titles = set(FAILS_SERIES_TO_TITLE.values())
        for column in table["columnsDefs"][1:]:
            assert column["cellDataType"] == "number"
            assert column["headerName"] == column["field"]
            if not column["field"].endswith(" Total"):
                assert column["field"] in model_titles

    def test_chart_data_types_make_the_series_chartable(self):
        """The date column is the time axis; every value column is a chart series."""
        columns = self._table()["columnsDefs"]
        assert columns[0]["field"] == "date"
        assert columns[0]["chartDataType"] == "time"
        value_columns = columns[1:]
        assert len(value_columns) == 14
        for column in value_columns:
            assert column["chartDataType"] == "series"


class TestPrimaryDealerPositioningWidget:
    """Tests for the ``primary_dealer_positioning`` widget column headers."""

    @staticmethod
    def _table() -> dict:
        """Return the ``primary_dealer_positioning`` widget table config."""
        from openbb_federal_reserve.regional import new_york

        route = next(
            r
            for r in new_york.router.api_router.routes
            if getattr(r, "name", "") == "primary_dealer_positioning"
        )
        return route.openapi_extra["widget_config"]["data"]["table"]  # ty: ignore[unresolved-attribute]

    def test_header_names_match_field_to_label_map(self):
        """Each non-date column carries the proper label from the field map."""
        from openbb_federal_reserve.utils.primary_dealer_statistics import (
            POSITION_FIELD_TO_LABEL,
        )

        table = self._table()
        assert table["showAll"] is True
        columns = table["columnsDefs"]
        assert columns[0]["field"] == "date"
        value_columns = columns[1:]
        assert [c["field"] for c in value_columns] == list(POSITION_FIELD_TO_LABEL)
        for column in value_columns:
            assert column["cellDataType"] == "number"
            assert column["formatterFn"] == "int"
            assert column["headerName"] == POSITION_FIELD_TO_LABEL[column["field"]]

    def test_label_map_covers_every_position_field(self):
        """The label map covers every ``dealer_position`` series field."""
        from openbb_federal_reserve.utils.primary_dealer_statistics import (
            POSITION_FIELD_TO_LABEL,
            POSITION_SERIES_TO_FIELD,
        )

        position_fields = set(POSITION_SERIES_TO_FIELD["dealer_position"].values())
        assert position_fields == set(POSITION_FIELD_TO_LABEL)
