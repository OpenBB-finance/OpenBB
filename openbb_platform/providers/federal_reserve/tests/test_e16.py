"""Tests for the FFIEC E.16 Country Exposure client and parser."""

import io
import zipfile
from datetime import date

from openpyxl import Workbook

from openbb_federal_reserve.utils import e16


class _FakeResponse:
    """Minimal stand-in for a curl_cffi response carrying text and bytes."""

    def __init__(self, text: str = "", content: bytes = b""):
        self.text = text
        self.content = content


class _FakeSession:
    """Path-aware fake returning the page HTML or the ZIP download."""

    def __init__(self, page: str = "", content: bytes = b""):
        self._page = page
        self._content = content
        self.urls: list[str] = []

    def get(self, url, timeout=None):
        """Return the listing page for the data page, else the ZIP content."""
        self.urls.append(url)
        if url == e16.PAGE_URL:
            return _FakeResponse(text=self._page)
        return _FakeResponse(content=self._content)


class TestGetSession:
    """Tests for the shared impersonating session delegate."""

    def test_delegates_to_shared_helper(self, monkeypatch):
        """``_get_session`` returns the shared session keyed by ``e16``."""
        sentinel = object()
        keys: list[str] = []
        monkeypatch.setattr(
            e16, "get_session", lambda key, warmup=None: keys.append(key) or sentinel
        )
        assert e16._get_session() is sentinel
        assert keys == ["e16"]


class TestFileQuarter:
    """Tests for the report-date to data-file quarter mapping."""

    def test_adds_three_months_with_year_rollover(self):
        """The data file is the report quarter end plus three months."""
        assert e16._file_quarter(date(2024, 3, 31)) == "202406"
        assert e16._file_quarter(date(2024, 12, 31)) == "202503"


class TestFetchE16:
    """Tests for the E.16 ZIP download flow."""

    def test_explicit_report_date_builds_url(self, monkeypatch):
        """A report date maps directly to its data-file quarter URL."""
        session = _FakeSession(content=b"PKzip")
        monkeypatch.setattr(e16, "_get_session", lambda: session)
        assert e16.fetch_e16(date(2024, 12, 31)) == b"PKzip"
        assert f"{e16.BASE_URL}/E16_202503.zip" in session.urls

    def test_latest_scrapes_page_for_quarter(self, monkeypatch):
        """With no report date the latest quarter is scraped from the page."""
        page = '<a href="/E16_202403.zip">x</a><a href="/E16_202412.zip">y</a>'
        session = _FakeSession(page=page, content=b"PKlatest")
        monkeypatch.setattr(e16, "_get_session", lambda: session)
        assert e16.fetch_e16() == b"PKlatest"
        assert f"{e16.BASE_URL}/E16_202412.zip" in session.urls

    def test_latest_no_quarters_returns_empty(self, monkeypatch):
        """A listing page with no data files returns empty bytes."""
        session = _FakeSession(page="<html></html>")
        monkeypatch.setattr(e16, "_get_session", lambda: session)
        assert e16.fetch_e16() == b""


def _zip(workbook: Workbook) -> bytes:
    """Pack a workbook into the cleansed-XLSX bulk ZIP layout."""
    book = io.BytesIO()
    workbook.save(book)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("E16 (009)_Cleansed.xlsx", book.getvalue())
    return buffer.getvalue()


class TestIsNumber:
    """Tests for the numeric-cell guard."""

    def test_numeric_and_non_numeric(self):
        """Numbers (with thousands separators) parse; text and None do not."""
        assert e16._is_number("1,234")
        assert e16._is_number(56)
        assert not e16._is_number("CANADA")
        assert not e16._is_number(None)


def _book_bytes(workbook: Workbook) -> bytes:
    """Serialize a workbook to bytes."""
    book = io.BytesIO()
    workbook.save(book)
    return book.getvalue()


class TestPeriod:
    """Tests for the cover-sheet period reader."""

    def test_unparseable_period_returns_none(self):
        """A malformed period date on the cover sheet resolves to ``None``."""
        from openpyxl import load_workbook

        workbook = Workbook()
        workbook.active["A1"] = "Period: Foo 99, 2024"
        loaded = load_workbook(io.BytesIO(_book_bytes(workbook)))
        assert e16._period(loaded) is None

    def test_no_period_marker_returns_none(self):
        """A cover sheet with no period marker resolves to ``None``."""
        from openpyxl import load_workbook

        workbook = Workbook()
        workbook.active["A1"] = "Statistical Release E.16"
        loaded = load_workbook(io.BytesIO(_book_bytes(workbook)))
        assert e16._period(loaded) is None


class TestParseSheet:
    """Tests for the long-format sheet parser's guards."""

    def test_no_xlsx_member_returns_empty(self):
        """Content with no XLSX member yields empty rows and no period."""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("readme.txt", "x")
        assert e16.parse_sheet(buffer.getvalue(), "All Banks - Table 1") == ([], None)

    def test_unknown_sheet_returns_period_only(self):
        """A sheet name absent from the workbook returns just the period."""
        workbook = Workbook()
        cover = workbook.active
        cover.title = "Cover"
        cover["A1"] = "Period: December 31, 2024"
        rows, period = e16.parse_sheet(_zip(workbook), "Missing Sheet")
        assert rows == []
        assert period == date(2024, 12, 31)

    def test_no_data_rows_returns_period_only(self):
        """A sheet with no numeric data rows returns just the period."""
        workbook = Workbook()
        cover = workbook.active
        cover.title = "Cover"
        cover["A1"] = "Period: December 31, 2024"
        sheet = workbook.create_sheet("All Banks - Table 1")
        sheet.cell(1, 2, "Header only, no numbers")
        rows, period = e16.parse_sheet(_zip(workbook), "All Banks - Table 1")
        assert rows == []
        assert period == date(2024, 12, 31)

    def test_blank_name_rows_are_skipped(self):
        """A data-area row with an empty name cell is skipped during the walk."""
        workbook = Workbook()
        cover = workbook.active
        cover.title = "Cover"
        cover["A1"] = "Period: December 31, 2024"
        sheet = workbook.create_sheet("All Banks - Table 1")
        sheet.cell(1, 3, "Claims")
        sheet.cell(2, 2, "CANADA")
        sheet.cell(2, 3, 179188)
        # A trailing row whose name cell is blank must be skipped.
        sheet.cell(3, 3, 5)
        rows, _ = e16.parse_sheet(_zip(workbook), "All Banks - Table 1")
        assert {row["country"] for row in rows} == {"CANADA"}

    def test_merged_headers_region_groups_and_labels(self):
        """Merged multi-level headers, region groups, and footnotes parse fully."""
        workbook = Workbook()
        cover = workbook.active
        cover.title = "Cover"
        cover["A1"] = "Period: December 31, 2024"
        sheet = workbook.create_sheet("All Banks - Table 1")
        sheet.cell(1, 2, "Country Exposure Survey Table 1")
        sheet.merge_cells("B1:F1")  # table-wide title, dropped from labels
        sheet.cell(2, 3, "Claims")
        sheet.merge_cells("C2:D2")
        sheet.cell(2, 5, "Off-Balance Sheet")
        sheet.merge_cells("E2:F2")
        sheet.cell(3, 3, "Cross-border Claims /2")  # footnote marker stripped
        sheet.cell(3, 4, "Total")
        sheet.cell(3, 5, "Unused Commitments")
        sheet.cell(3, 6, "Guarantees")
        sheet.cell(4, 2, "G-10 and Luxembourg")  # region header (no values)
        for row, country, values in (
            (5, "CANADA", (179188, 82747, 78777, 28438)),
            (6, "FRANCE", (152588, 20399, 220289, 6169)),
        ):
            sheet.cell(row, 2, country)
            for offset, value in enumerate(values):
                sheet.cell(row, 3 + offset, value)
        rows, period = e16.parse_sheet(_zip(workbook), "All Banks - Table 1")
        assert period == date(2024, 12, 31)
        assert all(r["country_group"] == "G-10 and Luxembourg" for r in rows)
        by_key = {(r["country"], r["label"]): r["value"] for r in rows}
        assert by_key[("CANADA", "Claims - Cross-border Claims")] == 179188 * 1_000_000
        assert (
            by_key[("CANADA", "Off-Balance Sheet - Unused Commitments")]
            == 78777 * 1_000_000
        )
