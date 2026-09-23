"""USDA ERS Livestock and Meat International Trade Data catalog, fetch, and parser."""

import io
import re

PRODUCT_PAGE = "data-products/livestock-and-meat-international-trade-data"

CARCASS_UNIT = "Carcass weight, 1,000 pounds"
HEAD_UNIT = "Head"

SPECIES: dict[str, dict] = {
    "beef_veal": {
        "label": "Beef and veal",
        "annual": "/media/5598/beef-and-veal-annual-and-cumulative-year-to-date"
        "-us-trade-carcass-weight-1000-pounds.xlsx",
        "monthly": "/media/5608/beef-and-veal-monthly-us-trade-carcass-weight"
        "-1000-pounds.xlsx",
        "unit": CARCASS_UNIT,
        "unit_in_label": False,
    },
    "cattle": {
        "label": "Cattle",
        "annual": "/media/5599/cattle-annual-and-cumulative-year-to-date"
        "-us-trade-head.xlsx",
        "monthly": "/media/5609/cattle-monthly-us-trade-head.xlsx",
        "unit": HEAD_UNIT,
        "unit_in_label": False,
    },
    "pork": {
        "label": "Pork",
        "annual": "/media/5600/pork-annual-and-cumulative-year-to-date"
        "-us-trade-carcass-weight-1000-pounds.xlsx",
        "monthly": "/media/5610/pork-monthly-us-trade-carcass-weight-1000-pounds.xlsx",
        "unit": CARCASS_UNIT,
        "unit_in_label": False,
    },
    "hogs": {
        "label": "Hogs",
        "annual": "/media/5601/hogs-annual-and-cumulative-year-to-date"
        "-us-trade-head.xlsx",
        "monthly": "/media/5611/hogs-monthly-us-trade-head.xlsx",
        "unit": HEAD_UNIT,
        "unit_in_label": False,
    },
    "lamb_mutton": {
        "label": "Lamb and mutton",
        "annual": "/media/5602/lamb-and-mutton-annual-and-cumulative-year-to-date"
        "-us-trade-carcass-weight-1000-pounds.xlsx",
        "monthly": "/media/5612/lamb-and-mutton-monthly-us-trade-carcass-weight"
        "-1000-pounds.xlsx",
        "unit": CARCASS_UNIT,
        "unit_in_label": False,
    },
    "sheep_goats": {
        "label": "Sheep and goats",
        "annual": "/media/5603/sheep-and-goats-annual-and-cumulative-year-to-date"
        "-us-trade-head.xlsx",
        "monthly": "/media/5613/sheep-and-goats-monthly-us-trade-head.xlsx",
        "unit": HEAD_UNIT,
        "unit_in_label": False,
    },
    "poultry_eggs": {
        "label": "Chickens, turkeys, and eggs",
        "annual": "/media/5604/chickens-turkeys-and-eggs-annual-and-cumulative"
        "-year-to-date-us-trade.xlsx",
        "monthly": "/media/5614/chickens-turkeys-and-eggs-monthly-us-trade.xlsx",
        "unit": None,
        "unit_in_label": True,
    },
}

SPECIES_LABELS: dict[str, str] = {key: cfg["label"] for key, cfg in SPECIES.items()}

DIRECTIONS: tuple[str, ...] = ("imports", "exports")

DIRECTION_LABELS: dict[str, str] = {"imports": "Imports", "exports": "Exports"}

FREQUENCIES: tuple[str, ...] = ("annual", "monthly")

FREQUENCY_LABELS: dict[str, str] = {"annual": "Annual", "monthly": "Monthly"}

MONTHS: tuple[str, ...] = (
    "Jan",
    "Feb",
    "Mar",
    "Apr",
    "May",
    "Jun",
    "Jul",
    "Aug",
    "Sep",
    "Oct",
    "Nov",
    "Dec",
)

MONTH_ORD: dict[str, int] = {name: index + 1 for index, name in enumerate(MONTHS)}

HEADER_MARKER = "Import/export"
FOOTER_PREFIXES = ("1/", "2/", "3/", "Source", "Date run", "USSR")
NULL_TOKENS = frozenset({"", "-", "--", "---", "n/a", "na", "null", "none", "nan"})

YEAR_CELL = re.compile(r"^\d{4}$")
MONTH_CELL = re.compile(r"^([A-Za-z]{3})-(\d{2})$")
YTD_CELL = re.compile(r"^([A-Za-z]{3}-[A-Za-z]{3})\s+(\d{2})$")
UNIT_PAREN = re.compile(r"\s*\(([^()]*)\)\s*$")


def normalize_label(value) -> str:
    """Collapse a source label's newlines and repeated spaces into one line.

    Parameters
    ----------
    value : object
        Raw column-A cell value.

    Returns
    -------
    str
        The label with newlines turned into spaces and whitespace collapsed.
    """
    return re.sub(r"\s+", " ", str(value).replace("\n", " ")).strip()


def expand_year(two_digit: int) -> int:
    """Expand a two-digit source year to a four-digit calendar year.

    Parameters
    ----------
    two_digit : int
        Two-digit year from a 'Mon-YY' or year-to-date header cell.

    Returns
    -------
    int
        The four-digit year, mapping 00-49 to the 2000s and 50-99 to the 1900s.
    """
    return 2000 + two_digit if two_digit < 50 else 1900 + two_digit


def parse_block_label(
    label: str, unit_in_label: bool, default_unit: str | None
) -> tuple[str, str, str | None]:
    """Split a block label into trade direction, product, and unit.

    Parameters
    ----------
    label : str
        Normalized block label, e.g. 'Cattle imports, total'.
    unit_in_label : bool
        Whether the trailing parenthetical carries the data unit, as for the
        chickens-turkeys-and-eggs workbook whose blocks mix meat and egg units.
    default_unit : str | None
        Unit to use when the label does not carry one.

    Returns
    -------
    tuple[str, str, str | None]
        The direction ('imports' or 'exports'), the product with the direction
        word removed, and the unit.
    """
    text = label
    unit = default_unit
    direction = "exports" if "export" in text.lower() else "imports"
    if unit_in_label:
        match = UNIT_PAREN.search(text)
        if match:
            unit = match.group(1).strip()
            text = text[: match.start()].strip()
    word = "exports" if direction == "exports" else "imports"
    product = re.sub(rf"\b{word}\b", "", text, count=1, flags=re.IGNORECASE)
    product = re.sub(r"\s+,", ",", product)
    product = re.sub(r"\s+", " ", product).strip(" ,")
    return direction, product, unit


def coerce_value(value) -> float | None:
    """Coerce a source cell to a float, keeping full precision.

    Parameters
    ----------
    value : object
        Raw value cell.

    Returns
    -------
    float | None
        The numeric value, or None for blank, placeholder, or non-numeric cells.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if text.casefold() in NULL_TOKENS:
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None


def parse_time_header(header: tuple) -> list[dict]:
    """Parse a header row's time cells into ordered column descriptors.

    Parameters
    ----------
    header : tuple
        The header row values; time cells begin at column D (index 3).

    Returns
    -------
    list[dict]
        One descriptor per recognized time cell with index, year, month,
        month_ord, col_key, and is_ytd keys.
    """
    columns: list[dict] = []
    for index in range(3, len(header)):
        cell = header[index]
        if cell is None:
            continue
        text = str(cell).strip()
        if YEAR_CELL.match(text):
            columns.append(
                {
                    "index": index,
                    "year": int(text),
                    "month": None,
                    "month_ord": 0,
                    "col_key": text,
                    "is_ytd": False,
                }
            )
            continue
        month_match = MONTH_CELL.match(text)
        if month_match:
            month = month_match.group(1).title()
            year = expand_year(int(month_match.group(2)))
            columns.append(
                {
                    "index": index,
                    "year": year,
                    "month": month,
                    "month_ord": MONTH_ORD.get(month, 0),
                    "col_key": str(year),
                    "is_ytd": False,
                }
            )
            continue
        ytd_match = YTD_CELL.match(text)
        if ytd_match:
            year = expand_year(int(ytd_match.group(2)))
            columns.append(
                {
                    "index": index,
                    "year": year,
                    "month": None,
                    "month_ord": 0,
                    "col_key": f"{ytd_match.group(1)} {year}",
                    "is_ytd": True,
                }
            )
    return columns


def _iter_sheets(content: bytes):
    """Yield each worksheet's row tuples from workbook bytes, recent sheet first.

    Parameters
    ----------
    content : bytes
        Raw XLSX bytes.

    Yields
    ------
    list[tuple]
        The value rows of one worksheet, iterating sheets in reverse order so
        the most recent split sheet establishes the country ranking.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(
        io.BytesIO(content), data_only=True, read_only=True
    )
    for sheet_name in reversed(workbook.sheetnames):
        yield list(workbook[sheet_name].iter_rows(values_only=True))
    workbook.close()


def _header_index(rows: list[tuple]) -> int:
    """Return the index of a sheet's 'Import/export' header row, or -1.

    Parameters
    ----------
    rows : list[tuple]
        A worksheet's value rows.

    Returns
    -------
    int
        The header row index, or -1 when the sheet has no header.
    """
    for index, row in enumerate(rows):
        cell = row[0] if row else None
        if cell is not None and str(cell).startswith(HEADER_MARKER):
            return index
    return -1


def parse_products(content: bytes, entry: dict, direction: str) -> list[str]:
    """List a direction's product blocks from an annual workbook, in file order.

    Parameters
    ----------
    content : bytes
        Raw XLSX bytes of a species annual workbook.
    entry : dict
        The species catalog entry from SPECIES.
    direction : str
        'imports' or 'exports'.

    Returns
    -------
    list[str]
        Distinct product labels for the direction, in source order.
    """
    products: list[str] = []
    seen: set[str] = set()
    for rows in _iter_sheets(content):
        start = _header_index(rows)
        if start < 0:
            continue
        for row in rows[start + 1 :]:
            label = row[0] if row else None
            code = row[1] if len(row) > 1 else None
            if label is None or code is None:
                continue
            block_direction, product, _ = parse_block_label(
                normalize_label(label), entry["unit_in_label"], entry["unit"]
            )
            if block_direction != direction or product in seen:
                continue
            seen.add(product)
            products.append(product)
    return products


def parse_workbook(
    content: bytes,
    entry: dict,
    want_direction: str,
    want_product: str,
) -> list[dict]:
    """Parse one species workbook into long records for a single product block.

    Parameters
    ----------
    content : bytes
        Raw XLSX bytes of the species annual or monthly workbook.
    entry : dict
        The species catalog entry from SPECIES.
    want_direction : str
        Trade direction to keep: 'imports' or 'exports'.
    want_product : str
        Product block to keep, matching parse_block_label's product.

    Returns
    -------
    list[dict]
        Records with direction, product, unit, geography_code, country, year,
        month, month_ord, col_key, is_ytd, rank, and value keys; only non-null
        observations are emitted, and 'Total' summary rows are dropped.
    """
    records: list[dict] = []
    ranks: dict[str, int] = {}
    for rows in _iter_sheets(content):
        start = _header_index(rows)
        if start < 0:
            continue
        columns = parse_time_header(rows[start])
        current: tuple[str, str, str | None] | None = None
        for row in rows[start + 1 :]:
            label = row[0] if row else None
            code = row[1] if len(row) > 1 else None
            name = row[2] if len(row) > 2 else None
            label_text = normalize_label(label) if label not in (None, "") else None
            country = str(name).strip() if name is not None else ""
            if label_text and code is None and not country:
                break
            if label_text and code is not None:
                current = parse_block_label(
                    label_text, entry["unit_in_label"], entry["unit"]
                )
            if current is None or not country:
                continue
            direction, product, unit = current
            if direction != want_direction or product != want_product:
                continue
            if country.casefold() == "total" or country.startswith("Total"):
                continue
            geography_code = (
                int(code)
                if isinstance(code, (int, float)) and not isinstance(code, bool)
                else None
            )
            if country not in ranks:
                ranks[country] = len(ranks)
            rank = ranks[country]
            for column in columns:
                value = coerce_value(
                    row[column["index"]] if column["index"] < len(row) else None
                )
                if value is None:
                    continue
                records.append(
                    {
                        "direction": direction,
                        "product": product,
                        "unit": unit,
                        "geography_code": geography_code,
                        "country": country,
                        "year": column["year"],
                        "month": column["month"],
                        "month_ord": column["month_ord"],
                        "col_key": column["col_key"],
                        "is_ytd": column["is_ytd"],
                        "rank": rank,
                        "value": value,
                    }
                )
    return records


async def afetch_products(table: str, direction: str) -> list[dict]:
    """Download the species annual workbook and list a direction's products.

    Parameters
    ----------
    table : str
        Species key from SPECIES.
    direction : str
        'imports' or 'exports'.

    Returns
    -------
    list[dict]
        Label/value option dictionaries for the product selector.
    """
    from openbb_government_us.usda.utils.ers_client import afetch_ers_file

    entry = SPECIES[table]
    content = await afetch_ers_file(entry["annual"], product=PRODUCT_PAGE)
    return [
        {"label": product, "value": product}
        for product in parse_products(content, entry, direction)
    ]


async def resolve_product(
    table: str, direction: str, product: str | None
) -> str | None:
    """Resolve a requested product to a valid one, defaulting to the first block.

    Parameters
    ----------
    table : str
        Species key from SPECIES.
    direction : str
        'imports' or 'exports'.
    product : str | None
        Requested product, or None to use the direction's first product.

    Returns
    -------
    str | None
        A product present in the workbook, or None when the direction has none.
    """
    options = await afetch_products(table, direction)
    values = [option["value"] for option in options]
    if product in values:
        return product
    return values[0] if values else None


async def afetch_records(
    table: str, direction: str, frequency: str, product: str | None
) -> list[dict]:
    """Download and parse one species-direction-product block for a frequency.

    Parameters
    ----------
    table : str
        Species key from SPECIES.
    direction : str
        'imports' or 'exports'.
    frequency : str
        'annual' or 'monthly'.
    product : str | None
        Product block, or None to use the direction's first product.

    Returns
    -------
    list[dict]
        Long records from parse_workbook for the resolved product block.
    """
    from openbb_government_us.usda.utils.ers_client import afetch_ers_file

    entry = SPECIES[table]
    resolved = await resolve_product(table, direction, product)
    if resolved is None:
        return []
    media_path = entry["monthly"] if frequency == "monthly" else entry["annual"]
    content = await afetch_ers_file(media_path, product=PRODUCT_PAGE)
    return parse_workbook(content, entry, direction, resolved)
