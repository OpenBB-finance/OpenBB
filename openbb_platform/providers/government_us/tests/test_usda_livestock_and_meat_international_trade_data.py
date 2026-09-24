"""Tests for the USDA ERS livestock and meat international trade data utils and model."""

import asyncio
import io

import openpyxl
import pytest
from openbb_core.app.model.abstract.error import OpenBBError
from openbb_core.provider.utils.errors import EmptyDataError
from pydantic.alias_generators import to_snake

from openbb_government_us.usda.models.livestock_and_meat_international_trade_data import (
    DIM_KEYS,
    LivestockAndMeatInternationalTradeDataData,
    LivestockAndMeatInternationalTradeDataFetcher,
    LivestockAndMeatInternationalTradeDataQueryParams,
)
from openbb_government_us.usda.utils import (
    ers_livestock_and_meat_international_trade_data as lmt,
)
from openbb_government_us.usda.utils.ers_livestock_and_meat_international_trade_data import (
    DIRECTIONS,
    FREQUENCIES,
    SPECIES,
    SPECIES_LABELS,
    coerce_value,
    expand_year,
    normalize_label,
    parse_block_label,
    parse_products,
    parse_time_header,
    parse_workbook,
)

HEAD_ENTRY = {"unit": "Head", "unit_in_label": False}

ANNUAL_HEADER = [
    "Import/export, geography code and name 1/ 2/",
    None,
    None,
    "1989",
    "1990",
    "Jan-May 25",
    "Jan-May 26",
]

ANNUAL_ROWS = [
    ["Cattle: annual and cumulative year-to-date U.S. trade (head)", None, None],
    ANNUAL_HEADER,
    ["Cattle imports, total", 1220, "Canada", 100.0, 110.0, 50.0, 55.0],
    [None, 2010, "Mexico", 200.0, None, 20.0, 25.0],
    [None, None, "Total\n ", 300.0, 110.0, 70.0, 80.0],
    ["Cattle exports, total", 5880, "Japan", 5.0, 6.0, 3.0, 4.0],
    [None, None, "Other geographies", 1.0, 2.0, None, 1.0],
    ["1/ Geographies are ranked by the sum of their trade.", None, None],
    ["Date run: 7/7/2026", None, None],
]

MONTHLY_SHEET0 = [
    ["Cattle: monthly U.S. trade (head)", None, None],
    ["Import/export, geography code and name 1/ 2/", None, None, "Jan-89", "Feb-89"],
    ["Cattle imports, total", 1220, "Canada", 11.0, 12.0],
    ["Cattle exports, total", 5880, "Japan", 1.0, 2.0],
    ["1/ ranked", None, None],
]

MONTHLY_SHEET1 = [
    ["Cattle: monthly U.S. trade (head)", None, None],
    ["Import/export, geography code and name 1/ 2/", None, None, "Jan-00", "Feb-26"],
    ["Cattle imports, total", 1220, "Canada", 21.0, 26.0],
    [None, 2010, "Mexico", 31.0, 36.0],
    ["Cattle exports, total", 5880, "Japan", 3.0, 4.0],
    ["1/ ranked", None, None],
]


def build_xlsx(sheets: list[tuple[str, list[list]]]) -> bytes:
    """Build workbook bytes from a list of (sheet name, rows) pairs."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        worksheet = workbook.create_sheet(title=name)
        for row in rows:
            worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


ANNUAL_BYTES = build_xlsx([("Cattle_Yearly", ANNUAL_ROWS)])
MONTHLY_BYTES = build_xlsx(
    [("Cattle_Monthly", MONTHLY_SHEET0), ("Cattle_Monthly(2)", MONTHLY_SHEET1)]
)


def make_record(**overrides) -> dict:
    """Build a parsed long record with optional field overrides."""
    record = {
        "direction": "imports",
        "product": "Cattle, total",
        "unit": "Head",
        "geography_code": 1220,
        "country": "Canada",
        "year": 2024,
        "month": None,
        "month_ord": 0,
        "col_key": "2024",
        "is_ytd": False,
        "rank": 0,
        "value": 100.0,
    }
    record.update(overrides)
    return record


class TestErsLivestockAndMeatInternationalTradeDataUtils:
    """Tests for the ers_livestock_and_meat_international_trade_data utils module."""

    def test_species_catalog(self):
        """The catalog holds the seven species, each with annual and monthly media."""
        assert len(SPECIES) == 7
        assert set(SPECIES) == set(SPECIES_LABELS)
        assert SPECIES["beef_veal"]["annual"].startswith("/media/5598/")
        assert SPECIES["beef_veal"]["monthly"].startswith("/media/5608/")
        assert SPECIES["poultry_eggs"]["unit_in_label"] is True
        assert SPECIES["cattle"]["unit"] == "Head"
        assert DIRECTIONS == ("imports", "exports")
        assert FREQUENCIES == ("annual", "monthly")

    def test_normalize_label(self):
        """Newlines and repeated spaces collapse to a single line."""
        assert normalize_label("Beef and veal\nimports") == "Beef and veal imports"
        assert normalize_label("Other chicken\nimports\n(1,000 pounds)") == (
            "Other chicken imports (1,000 pounds)"
        )

    def test_expand_year(self):
        """Two-digit years map to the 1900s at or above 50 and the 2000s below."""
        assert expand_year(89) == 1989
        assert expand_year(99) == 1999
        assert expand_year(0) == 2000
        assert expand_year(26) == 2026
        assert expand_year(49) == 2049
        assert expand_year(50) == 1950

    def test_parse_block_label_direction_at_end(self):
        """A trailing direction word yields the leading product and file unit."""
        assert parse_block_label("Beef and veal imports", False, "Head") == (
            "imports",
            "Beef and veal",
            "Head",
        )
        assert parse_block_label("Mutton exports", False, "Head") == (
            "exports",
            "Mutton",
            "Head",
        )

    def test_parse_block_label_subcategory(self):
        """A species-direction-subcategory label keeps the subcategory."""
        assert parse_block_label("Cattle imports, total", False, "Head") == (
            "imports",
            "Cattle, total",
            "Head",
        )
        assert parse_block_label(
            "Hog imports, less than 7 kilograms (15.4 pounds)", False, "Head"
        ) == ("imports", "Hog, less than 7 kilograms (15.4 pounds)", "Head")

    def test_parse_block_label_unit_in_label(self):
        """The chickens-turkeys-and-eggs blocks pull the unit from the label."""
        assert parse_block_label("Broiler imports (1,000 pounds)", True, None) == (
            "imports",
            "Broiler",
            "1,000 pounds",
        )
        assert parse_block_label(
            "Egg imports, including products (shell-egg equivalent, 1,000 dozen)",
            True,
            None,
        ) == ("imports", "Egg, including products", "shell-egg equivalent, 1,000 dozen")

    def test_coerce_value(self):
        """Values keep full precision; blanks, placeholders, and text become None."""
        assert coerce_value(818411.763373013) == 818411.763373013
        assert coerce_value(5) == 5.0
        assert coerce_value("1,234.5") == 1234.5
        assert coerce_value(None) is None
        assert coerce_value(True) is None
        assert coerce_value("--") is None
        assert coerce_value("n/a") is None
        assert coerce_value("abc") is None

    def test_parse_time_header_annual(self):
        """Annual headers yield year columns then the trailing year-to-date columns."""
        columns = parse_time_header(tuple(ANNUAL_HEADER))
        assert [c["col_key"] for c in columns] == [
            "1989",
            "1990",
            "Jan-May 2025",
            "Jan-May 2026",
        ]
        assert [c["is_ytd"] for c in columns] == [False, False, True, True]
        assert [c["year"] for c in columns] == [1989, 1990, 2025, 2026]
        assert all(c["month"] is None for c in columns)

    def test_parse_time_header_monthly_and_unknown(self):
        """Monthly headers split into year and month; unknown cells are skipped."""
        header = (None, None, None, "Jan-89", "May-26", "junk", None)
        columns = parse_time_header(header)
        assert [(c["col_key"], c["month"], c["month_ord"]) for c in columns] == [
            ("1989", "Jan", 1),
            ("2026", "May", 5),
        ]

    def test_parse_products(self):
        """Products list per direction in source order, deduplicated."""
        assert parse_products(ANNUAL_BYTES, HEAD_ENTRY, "imports") == ["Cattle, total"]
        assert parse_products(ANNUAL_BYTES, HEAD_ENTRY, "exports") == ["Cattle, total"]

    def test_parse_products_no_header(self):
        """A sheet without the header row contributes no products."""
        content = build_xlsx([("Empty", [["nothing here", None, None]])])
        assert parse_products(content, HEAD_ENTRY, "imports") == []

    def test_parse_workbook_annual_drops_total_keeps_other(self):
        """Annual parsing drops 'Total' rows and keeps 'Other geographies'."""
        records = parse_workbook(ANNUAL_BYTES, HEAD_ENTRY, "imports", "Cattle, total")
        countries = {r["country"] for r in records}
        assert countries == {"Canada", "Mexico"}
        canada = {r["col_key"]: r["value"] for r in records if r["country"] == "Canada"}
        assert canada == {
            "1989": 100.0,
            "1990": 110.0,
            "Jan-May 2025": 50.0,
            "Jan-May 2026": 55.0,
        }
        mexico = {r["col_key"]: r["value"] for r in records if r["country"] == "Mexico"}
        assert "1990" not in mexico
        assert all(r["unit"] == "Head" for r in records)
        assert all(r["month"] is None for r in records)

    def test_parse_workbook_other_geographies_null_code(self):
        """The 'Other geographies' residual is kept with a null geography code."""
        records = parse_workbook(ANNUAL_BYTES, HEAD_ENTRY, "exports", "Cattle, total")
        other = [r for r in records if r["country"] == "Other geographies"]
        assert other
        assert all(r["geography_code"] is None for r in other)
        japan = [r for r in records if r["country"] == "Japan"]
        assert all(r["geography_code"] == 5880 for r in japan)

    def test_parse_workbook_direction_and_product_filter(self):
        """Only the requested direction and product block is emitted."""
        records = parse_workbook(ANNUAL_BYTES, HEAD_ENTRY, "exports", "Cattle, total")
        assert {r["direction"] for r in records} == {"exports"}
        assert {r["product"] for r in records} == {"Cattle, total"}
        assert parse_workbook(ANNUAL_BYTES, HEAD_ENTRY, "imports", "Missing") == []

    def test_parse_workbook_monthly_merges_sheets(self):
        """Monthly parsing merges both sheets and populates the month dimension."""
        records = parse_workbook(MONTHLY_BYTES, HEAD_ENTRY, "imports", "Cattle, total")
        keys = {(r["country"], r["month"], r["col_key"]): r["value"] for r in records}
        assert keys[("Canada", "Jan", "1989")] == 11.0
        assert keys[("Canada", "Feb", "1989")] == 12.0
        assert keys[("Canada", "Jan", "2000")] == 21.0
        assert keys[("Canada", "Feb", "2026")] == 26.0
        assert keys[("Mexico", "Jan", "2000")] == 31.0
        assert {r["month"] for r in records} == {"Jan", "Feb"}

    def test_parse_workbook_monthly_rank_recent_first(self):
        """The recent split sheet establishes the country ranking order."""
        records = parse_workbook(MONTHLY_BYTES, HEAD_ENTRY, "imports", "Cattle, total")
        ranks = {r["country"]: r["rank"] for r in records}
        assert ranks["Canada"] == 0
        assert ranks["Mexico"] == 1

    def test_parse_workbook_skips_headerless_and_preblock_rows(self):
        """Headerless sheets and rows before any block label are skipped."""
        content = build_xlsx(
            [
                ("NoHeader", [["random", None, None], ["stuff", 1, "X"]]),
                (
                    "Cattle_Yearly",
                    [
                        ["title", None, None],
                        ANNUAL_HEADER,
                        [None, 9999, "Ghost", 1.0, 2.0, 3.0, 4.0],
                        [
                            "Cattle imports, total",
                            1220,
                            "Canada",
                            100.0,
                            110.0,
                            50.0,
                            55.0,
                        ],
                        ["Date run: 7/7/2026", None, None],
                    ],
                ),
            ]
        )
        records = parse_workbook(content, HEAD_ENTRY, "imports", "Cattle, total")
        assert {r["country"] for r in records} == {"Canada"}

    def test_afetch_products(self, monkeypatch):
        """afetch_products downloads the annual workbook and lists options."""

        async def fake_afetch_ers_file(media_path, product=None, ttl=None):
            assert media_path == SPECIES["cattle"]["annual"]
            return ANNUAL_BYTES

        monkeypatch.setattr(
            "openbb_government_us.usda.utils.ers_client.afetch_ers_file",
            fake_afetch_ers_file,
        )
        options = asyncio.run(lmt.afetch_products("cattle", "imports"))
        assert options == [{"label": "Cattle, total", "value": "Cattle, total"}]

    def test_resolve_product_valid_default_and_missing(self, monkeypatch):
        """resolve_product keeps a valid product and defaults to the first block."""

        async def fake_products(table, direction):
            return [{"label": "A", "value": "A"}, {"label": "B", "value": "B"}]

        monkeypatch.setattr(lmt, "afetch_products", fake_products)
        assert asyncio.run(lmt.resolve_product("cattle", "imports", "B")) == "B"
        assert asyncio.run(lmt.resolve_product("cattle", "imports", None)) == "A"
        assert asyncio.run(lmt.resolve_product("cattle", "imports", "Z")) == "A"

    def test_resolve_product_empty(self, monkeypatch):
        """resolve_product returns None when the direction has no products."""

        async def fake_products(table, direction):
            return []

        monkeypatch.setattr(lmt, "afetch_products", fake_products)
        assert asyncio.run(lmt.resolve_product("cattle", "imports", None)) is None

    def test_afetch_records_annual(self, monkeypatch):
        """afetch_records resolves the product and parses the annual workbook."""

        async def fake_resolve(table, direction, product):
            return "Cattle, total"

        async def fake_afetch_ers_file(media_path, product=None, ttl=None):
            assert media_path == SPECIES["cattle"]["annual"]
            return ANNUAL_BYTES

        monkeypatch.setattr(lmt, "resolve_product", fake_resolve)
        monkeypatch.setattr(
            "openbb_government_us.usda.utils.ers_client.afetch_ers_file",
            fake_afetch_ers_file,
        )
        records = asyncio.run(lmt.afetch_records("cattle", "imports", "annual", None))
        assert {r["country"] for r in records} == {"Canada", "Mexico"}

    def test_afetch_records_monthly_media(self, monkeypatch):
        """afetch_records selects the monthly workbook for the monthly frequency."""

        async def fake_resolve(table, direction, product):
            return "Cattle, total"

        seen = {}

        async def fake_afetch_ers_file(media_path, product=None, ttl=None):
            seen["media"] = media_path
            return MONTHLY_BYTES

        monkeypatch.setattr(lmt, "resolve_product", fake_resolve)
        monkeypatch.setattr(
            "openbb_government_us.usda.utils.ers_client.afetch_ers_file",
            fake_afetch_ers_file,
        )
        records = asyncio.run(
            lmt.afetch_records("cattle", "imports", "monthly", "Cattle, total")
        )
        assert seen["media"] == SPECIES["cattle"]["monthly"]
        assert {r["month"] for r in records} == {"Jan", "Feb"}

    def test_afetch_records_unresolved_product(self, monkeypatch):
        """afetch_records returns no records when no product resolves."""

        async def fake_resolve(table, direction, product):
            return None

        monkeypatch.setattr(lmt, "resolve_product", fake_resolve)
        assert (
            asyncio.run(lmt.afetch_records("cattle", "imports", "annual", None)) == []
        )


class TestLivestockAndMeatInternationalTradeData:
    """Tests for the LivestockAndMeatInternationalTradeData model."""

    def test_transform_query(self):
        """transform_query builds the query params model."""
        query = LivestockAndMeatInternationalTradeDataFetcher.transform_query(
            {"table": "pork", "direction": "exports", "frequency": "monthly"}
        )
        assert isinstance(query, LivestockAndMeatInternationalTradeDataQueryParams)
        assert query.table == "pork"
        assert query.direction == "exports"
        assert query.frequency == "monthly"

    def test_table_defaults_and_blank(self):
        """Table defaults to beef_veal and blanks or lists fall back correctly."""
        assert LivestockAndMeatInternationalTradeDataQueryParams().table == "beef_veal"
        assert (
            LivestockAndMeatInternationalTradeDataQueryParams(table="").table
            == "beef_veal"
        )
        assert (
            LivestockAndMeatInternationalTradeDataQueryParams(table=["cattle"]).table
            == "cattle"
        )

    def test_unknown_table_raises(self):
        """Unknown tables raise OpenBBError listing valid choices."""
        with pytest.raises(OpenBBError, match="Invalid table: bogus.*beef_veal"):
            LivestockAndMeatInternationalTradeDataQueryParams(table="bogus")

    def test_direction_defaults_and_normalizes(self):
        """Direction defaults to imports, lowercases, and takes a list's first."""
        assert (
            LivestockAndMeatInternationalTradeDataQueryParams().direction == "imports"
        )
        assert (
            LivestockAndMeatInternationalTradeDataQueryParams(direction="").direction
            == "imports"
        )
        assert (
            LivestockAndMeatInternationalTradeDataQueryParams(
                direction="EXPORTS"
            ).direction
            == "exports"
        )
        assert (
            LivestockAndMeatInternationalTradeDataQueryParams(
                direction=["exports"]
            ).direction
            == "exports"
        )

    def test_unknown_direction_raises(self):
        """An unknown direction raises OpenBBError."""
        with pytest.raises(OpenBBError, match="Invalid direction: sideways"):
            LivestockAndMeatInternationalTradeDataQueryParams(direction="sideways")

    def test_frequency_defaults_and_normalizes(self):
        """Frequency defaults to annual, lowercases, and takes a list's first."""
        assert LivestockAndMeatInternationalTradeDataQueryParams().frequency == "annual"
        assert (
            LivestockAndMeatInternationalTradeDataQueryParams(frequency="").frequency
            == "annual"
        )
        assert (
            LivestockAndMeatInternationalTradeDataQueryParams(
                frequency="MONTHLY"
            ).frequency
            == "monthly"
        )
        assert (
            LivestockAndMeatInternationalTradeDataQueryParams(
                frequency=["monthly"]
            ).frequency
            == "monthly"
        )

    def test_unknown_frequency_raises(self):
        """An unknown frequency raises OpenBBError."""
        with pytest.raises(OpenBBError, match="Invalid frequency: hourly"):
            LivestockAndMeatInternationalTradeDataQueryParams(frequency="hourly")

    def test_product_normalizes_to_none(self):
        """Product normalizes blanks to None and takes a list's first."""
        assert LivestockAndMeatInternationalTradeDataQueryParams().product is None
        assert (
            LivestockAndMeatInternationalTradeDataQueryParams(product="").product
            is None
        )
        assert (
            LivestockAndMeatInternationalTradeDataQueryParams(
                product=["Shell-egg"]
            ).product
            == "Shell-egg"
        )

    def test_aextract_data_forwards_params(self, monkeypatch):
        """aextract_data forwards every selection to afetch_records."""
        calls = []

        async def fake_afetch_records(table, direction, frequency, product):
            calls.append((table, direction, frequency, product))
            return [make_record()]

        monkeypatch.setattr(lmt, "afetch_records", fake_afetch_records)
        query = LivestockAndMeatInternationalTradeDataFetcher.transform_query(
            {"table": "cattle", "direction": "exports", "frequency": "monthly"}
        )
        records = asyncio.run(
            LivestockAndMeatInternationalTradeDataFetcher.aextract_data(query, None)
        )
        assert calls == [("cattle", "exports", "monthly", None)]
        assert records == [make_record()]

    def test_transform_data_pivots_years_into_columns(self):
        """Each country becomes one row with a value column per year."""
        records = [
            make_record(col_key="1989", year=1989, value=100.0),
            make_record(col_key="1990", year=1990, value=110.0),
            make_record(col_key="Jan-May 2026", year=2026, is_ytd=True, value=55.0),
        ]
        query = LivestockAndMeatInternationalTradeDataFetcher.transform_query({})
        data = LivestockAndMeatInternationalTradeDataFetcher.transform_data(
            query, records
        )
        assert len(data) == 1
        dumped = data[0].model_dump(by_alias=True)
        assert dumped["country"] == "Canada"
        assert dumped["1989"] == 100.0
        assert dumped["Jan-May 2026"] == 55.0
        assert not any(key.startswith("_") for key in dumped)

    def test_transform_data_year_columns_chronological_ytd_trailing(self):
        """Year columns emit oldest-first with year-to-date columns trailing."""
        records = [
            make_record(col_key="Jan-May 2026", year=2026, is_ytd=True, value=1.0),
            make_record(col_key="2025", year=2025, value=2.0),
            make_record(col_key="Jan-May 2025", year=2025, is_ytd=True, value=3.0),
            make_record(col_key="1989", year=1989, value=4.0),
        ]
        query = LivestockAndMeatInternationalTradeDataFetcher.transform_query({})
        data = LivestockAndMeatInternationalTradeDataFetcher.transform_data(
            query, records
        )
        dumped = data[0].model_dump(by_alias=True)
        dynamic = [k for k in dumped if k not in DIM_KEYS]
        assert dynamic == ["1989", "2025", "Jan-May 2025", "Jan-May 2026"]

    def test_transform_data_monthly_rows(self):
        """Monthly records fold the month into the partner label, sorted by rank."""
        records = [
            make_record(country="Japan", month="Feb", month_ord=2, rank=0, value=2.0),
            make_record(country="Japan", month="Jan", month_ord=1, rank=0, value=1.0),
            make_record(country="Mexico", month="Jan", month_ord=1, rank=1, value=3.0),
        ]
        query = LivestockAndMeatInternationalTradeDataFetcher.transform_query(
            {"frequency": "monthly"}
        )
        data = LivestockAndMeatInternationalTradeDataFetcher.transform_data(
            query, records
        )
        assert [(row.country, row.month) for row in data] == [
            ("Japan — Jan", "Jan"),
            ("Japan — Feb", "Feb"),
            ("Mexico — Jan", "Jan"),
        ]

    def test_transform_data_filters_years(self):
        """start_year and end_year select which year columns appear."""
        records = [
            make_record(col_key=str(year), year=year, value=float(year))
            for year in (2022, 2023, 2024)
        ]
        query = LivestockAndMeatInternationalTradeDataFetcher.transform_query(
            {"start_year": 2023, "end_year": 2023}
        )
        data = LivestockAndMeatInternationalTradeDataFetcher.transform_data(
            query, records
        )
        dumped = data[0].model_dump(by_alias=True)
        assert [k for k in dumped if k not in DIM_KEYS] == ["2023"]

    def test_transform_data_empty_raises(self):
        """No records surviving the filters raises EmptyDataError."""
        query = LivestockAndMeatInternationalTradeDataFetcher.transform_query(
            {"start_year": 2100}
        )
        with pytest.raises(EmptyDataError, match="No records match"):
            LivestockAndMeatInternationalTradeDataFetcher.transform_data(
                query, [make_record(year=2024)]
            )

    def test_columns_defs_bind_to_served_keys(self):
        """Every column definition binds to a served key; year columns vary."""
        records = [
            make_record(country="Canada", col_key="1989", year=1989, value=100.0),
            make_record(country="Canada", col_key="1990", year=1990, value=110.0),
            make_record(
                country="Mexico",
                geography_code=2010,
                rank=1,
                col_key="1989",
                year=1989,
                value=5.0,
            ),
            make_record(
                country="Mexico",
                geography_code=2010,
                rank=1,
                col_key="1990",
                year=1990,
                value=6.0,
            ),
        ]
        query = LivestockAndMeatInternationalTradeDataFetcher.transform_query({})
        data = LivestockAndMeatInternationalTradeDataFetcher.transform_data(
            query, records
        )
        dumped = [row.model_dump(by_alias=True) for row in data]
        served = set(dumped[0])
        schema = LivestockAndMeatInternationalTradeDataData.model_json_schema(
            by_alias=True, mode="serialization"
        )
        column_fields = [to_snake(key) for key in schema["properties"]]
        assert column_fields, "no column definitions generated"
        for field in column_fields:
            assert field in served, field
        assert set(column_fields) == set(DIM_KEYS)
        dynamic = [key for key in served if key not in column_fields]
        for key in dynamic:
            values = {row.get(key) for row in dumped}
            assert len(values) > 1, f"dynamic column {key} is constant"

    def test_widget_config_metadata(self):
        """The whole-widget config carries the ERS category and source."""
        config = LivestockAndMeatInternationalTradeDataData.model_config[
            "json_schema_extra"
        ]["x-widget_config"]
        assert config["$.category"] == "Economy"
        assert config["$.subCategory"] == "Agriculture"
        assert config["$.source"] == ["USDA", "ERS"]
        assert config["$.name"].startswith("USDA ERS")
